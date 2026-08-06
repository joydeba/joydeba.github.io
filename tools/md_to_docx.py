#!/usr/bin/env python3
import sys
from pathlib import Path
from docx import Document

def md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    text = md_path.read_text(encoding='utf-8')
    for line in text.splitlines():
        line = line.rstrip()
        if not line:
            doc.add_paragraph('')
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(line[2:].strip())
        elif line.startswith('+ '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(line[2:].strip())
        else:
            doc.add_paragraph(line)
    doc.save(str(docx_path))

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: md_to_docx.py input.md output.docx', file=sys.stderr)
        sys.exit(2)
    md_to_docx(Path(sys.argv[1]), Path(sys.argv[2]))

