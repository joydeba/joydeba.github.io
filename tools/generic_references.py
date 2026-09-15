#!/usr/bin/env python3
"""Generate a posting-agnostic academic References document (DOCX + PDF).

Styled to match Debasish Chakroborti academic CV (amber/green accents).
Output is written to a gitignored folder for upload to application portals.
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "Generic-Application-Docs"

AMBER = RGBColor(0xB9, 0x77, 0x0E)
GREEN = RGBColor(0x1E, 0x84, 0x49)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x44, 0x44, 0x44)

NAME = "Debasish Chakroborti (Joy), Ph.D."
EMAIL = "debasish.chakroborti@usask.ca"
PHONE = "306 914 6212"
ADDR = "410 Stensrud Rd #118, Saskatoon, SK"
WEB = "https://joydeba.github.io/"
ORCID = "https://orcid.org/0000-0002-1597-8162"

REFEREES = [
    {
        "name": "Professor Chanchal K. Roy",
        "tex_name": r"Professor Chanchal K.\ Roy",
        "rank": "Professor of Software Engineering",
        "dept": "Department of Computer Science",
        "university": "University of Saskatchewan",
        "street": "176 Thorvaldson Building, 110 Science Place",
        "city": "Saskatoon, SK  S7N 5C9, Canada",
        "email": "chanchal.roy@usask.ca",
        "phone": "+1 306 715 0600",
        "note": (
            "Professor Roy was co-supervisor of my M.Sc.\\ and Ph.D.\\ research and is my "
            "current postdoctoral supervisor in the Software Research Lab. He is able to "
            "comment on my research independence, scholarly publication, and collaborative "
            "research with industry in software analytics."
        ),
        "note_plain": (
            "Professor Roy was co-supervisor of my M.Sc. and Ph.D. research and is my "
            "current postdoctoral supervisor in the Software Research Lab. He is able to "
            "comment on my research independence, scholarly publication, and collaborative "
            "research with industry in software analytics."
        ),
    },
    {
        "name": "Professor Kevin A. Schneider",
        "tex_name": r"Professor Kevin A.\ Schneider",
        "rank": "Professor of Computer Science",
        "dept": "Department of Computer Science",
        "university": "University of Saskatchewan",
        "street": "176 Thorvaldson Building, 110 Science Place",
        "city": "Saskatoon, SK  S7N 5C9, Canada",
        "email": "kevin.schneider@usask.ca",
        "phone": "+1 306 966 4891",
        "note": (
            "Professor Schneider was co-supervisor of my M.Sc.\\ and Ph.D.\\ research. He is "
            "able to comment on the quality of my research, my software engineering "
            "work, and my training of graduate students."
        ),
        "note_plain": (
            "Professor Schneider was co-supervisor of my M.Sc. and Ph.D. research. He is "
            "able to comment on the quality of my research, my software engineering "
            "work, and my training of graduate students."
        ),
    },
    {
        "name": "Dr. Banani Roy",
        "tex_name": r"Dr.\ Banani Roy",
        "rank": "Associate Professor",
        "dept": "Department of Computer Science",
        "university": "University of Saskatchewan",
        "street": "176 Thorvaldson Building, 110 Science Place",
        "city": "Saskatoon, SK  S7N 5C9, Canada",
        "email": "banani.roy@usask.ca",
        "phone": "+1 306 850 5630",
        "note": (
            "Dr.\\ Roy was my M.Sc.\\ supervisor on scientific workflow research in the "
            "Plant Phenotyping and Imaging Research Centre (P2IRC) and served on my "
            "Ph.D.\\ advisory committee. We have co-authored "
            "publications. She is able to comment on my applied artificial intelligence "
            "systems research, interdisciplinary collaboration, and mentoring."
        ),
        "note_plain": (
            "Dr. Roy was my M.Sc. supervisor on scientific workflow research in the "
            "Plant Phenotyping and Imaging Research Centre (P2IRC) and served on my "
            "Ph.D. advisory committee. We have co-authored "
            "publications. She is able to comment on my applied artificial intelligence "
            "systems research, interdisciplinary collaboration, and mentoring."
        ),
    },
]


def set_run(run, *, size=10, bold=False, italic=False, color=DARK, font="Georgia"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_horizontal_line(paragraph, color_hex="B9770E"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hyperlink(paragraph, text, url, *, size=10, color=GREEN, font="Georgia"):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    r_font = OxmlElement("w:rFonts")
    r_font.set(qn("w:ascii"), font)
    r_font.set(qn("w:hAnsi"), font)
    rPr.append(r_font)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), str(color))
    rPr.append(color_el)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def setup_doc(margins=0.85):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(margins)
        section.bottom_margin = Inches(margins)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return doc


def header_block(doc, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(NAME)
    set_run(r, size=16, bold=True, color=DARK, font="Calibri")

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(subtitle)
        set_run(r2, size=11, bold=True, color=AMBER, font="Calibri")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    r = contact.add_run(f"{ADDR}  ·  {PHONE}  ·  {EMAIL}")
    set_run(r, size=9, color=MUTED, font="Calibri")

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(6)
    r = links.add_run(f"{WEB}  ·  {ORCID}")
    set_run(r, size=9, color=GREEN, font="Calibri")
    add_horizontal_line(links)


def meta_line(doc, text):
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(text)
    set_run(r, size=9.5, italic=True, color=MUTED, font="Calibri")
    return meta


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, size=11, bold=True, color=AMBER, font="Calibri")
    add_horizontal_line(p, "B9770E")
    return p


def body_para(doc, text, *, space_after=2, size=10, italic=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    set_run(r, size=size, italic=italic, bold=bold, color=DARK, font="Georgia")
    return p


def labeled_line(
    doc, label, text, *, space_after=2, size=10, email=None, phone=None, phone_label="Telephone: "
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r0 = p.add_run(label)
    set_run(r0, size=size, bold=True, color=DARK, font="Georgia")
    if email is not None and phone is not None:
        add_hyperlink(p, email, f"mailto:{email}", size=size, color=GREEN)
        r = p.add_run(f"    {phone_label}{phone}")
        set_run(r, size=size, color=DARK, font="Georgia")
    else:
        r = p.add_run(text)
        set_run(r, size=size, color=DARK, font="Georgia")
    return p


def write_docx():
    doc = setup_doc(0.85)
    header_block(doc, "List of Academic References")
    meta_line(
        doc,
        "The names and complete contact information of three academic referees are provided below.",
    )

    for ref in REFEREES:
        section_heading(doc, ref["name"])
        body_para(doc, ref["rank"], space_after=1)
        body_para(doc, ref["dept"], space_after=1)
        body_para(doc, ref["university"], space_after=1)
        body_para(doc, ref["street"], space_after=1)
        body_para(doc, ref["city"], space_after=3)
        labeled_line(
            doc,
            "Email: ",
            "",
            space_after=4,
            email=ref["email"],
            phone=ref["phone"],
            phone_label="Telephone: ",
        )
        body_para(doc, ref["note_plain"], space_after=8, size=10)

    path = OUT / "References.docx"
    doc.save(path)
    return path


def write_tex() -> Path:
    blocks = []
    for ref in REFEREES:
        blocks.append(
            "\\sectionheader{" + ref["tex_name"] + "}\n\n"
            "{\\setlength{\\parskip}{0.08em}\\setstretch{1.02}%\n"
            + ref["rank"] + "\\\\\n"
            + ref["dept"] + "\\\\\n"
            + ref["university"] + "\\\\\n"
            + ref["street"] + "\\\\\n"
            + ref["city"] + "\\par}\n"
            "\\vspace{0.32em}\n"
            "Email: \\href{mailto:" + ref["email"] + "}{" + ref["email"] + "}"
            "\\hspace{1.5em} Telephone: " + ref["phone"] + "\\par\n"
            "\\vspace{0.18em}\n"
            + ref["note"] + "\n"
        )
    body = "\n".join(blocks)
    tex = r"""\documentclass[10pt,english]{article}
\usepackage[letterpaper,nohead,nofoot,hmargin=0.85in,vmargin=0.75in]{geometry}
\usepackage[T1]{fontenc}
\usepackage[english]{babel}
\usepackage[p,osf,swashQ]{cochineal}
\usepackage{cabin}
\usepackage[varqu,varl,scale=0.9]{zi4}
\usepackage{setspace}
\usepackage{xcolor}
\usepackage{tikz}
\usepackage[fixed]{fontawesome5}
\usepackage{marvosym}
\usepackage[colorlinks=true,allcolors=black,breaklinks=true]{hyperref}

\definecolor{SwishLineColour}{HTML}{b9770e}
\definecolor{MarkerColour}{HTML}{1e8449}

\pagestyle{empty}
\setlength{\parskip}{0.42em}
\setlength{\parindent}{0pt}
\setstretch{1.06}

\newcommand{\makefield}[2]{%
  \mbox{\makebox[1.25em][l]{\color{MarkerColour!80!black}#1}\hspace{0.45em}#2}%
  \hspace{0.7em}\allowbreak
  \ignorespaces
}

\newcommand{\packetheader}[1]{%
{\LARGE\bfseries\sffamily Debasish Chakroborti (Joy), Ph.D.\par}
\vspace{0.2em}
{\raggedright
\makefield{\faGlobe}{\url{https://joydeba.github.io/}}%
\makefield{\faEnvelope}{\texttt{debasish.chakroborti@usask.ca}}%
\makefield{\faHome}{\texttt{410 Stensrud Rd \#118, Saskatoon, SK}}%
\makefield{\Mobilefone}{\texttt{306 914 6212}}%
\par}
\vspace{0.28em}
{\Large\bfseries\sffamily #1\par}
\vspace{0.18em}
\noindent\begin{tikzpicture}[baseline]
  \shade[left color=SwishLineColour!60!white, right color=white] rectangle (\linewidth,1.5pt);
\end{tikzpicture}
\vspace{0.38em}
}

\newcommand{\sectionheader}[1]{%
  \par\vspace{0.72em}%
  \noindent\begin{tikzpicture}[baseline]
  \shade[left color=SwishLineColour!60!white, right color=white] rectangle (\linewidth,1.5pt);
  \node[font={\large\bfseries\sffamily},inner sep=0pt,anchor=south west,text depth=.5ex,text height=1.5ex] at (1pt,2pt) {#1};
  \end{tikzpicture}%
  \par\vspace{0.32em}%
}

\begin{document}
\packetheader{List of Academic References}

{\centering\textit{The names and complete contact information of three academic referees are provided below.}\par}

<<BODY>>
\end{document}
""".replace("<<BODY>>", body)
    path = OUT / "References.tex"
    path.write_text(tex, encoding="utf-8")
    return path


def compile_pdf(tex_path: Path) -> Path:
    subprocess.run(
        ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex_path.name],
        cwd=tex_path.parent,
        check=True,
        capture_output=True,
        text=True,
    )
    for ext in (".aux", ".log", ".out"):
        extra = tex_path.with_suffix(ext)
        if extra.exists():
            extra.unlink()
    return tex_path.with_suffix(".pdf")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    docx_path = write_docx()
    tex_path = write_tex()
    pdf_path = compile_pdf(tex_path)
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
