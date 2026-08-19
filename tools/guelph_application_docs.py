#!/usr/bin/env python3
"""Generate University of Guelph (Req. 2649) application DOCX files
styled to match Debasish Chakroborti academic CV (amber/green accents).

Position: Assistant Professor, Software Engineering & AI / Applied AI
School of Computer Science, College of Computational Mathematical and
Physical Sciences. Earliest start January 1, 2027. Review begins Aug 10, 2026.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parents[1] / "University-of-Guelph-SoCS-2649"

# Match cv-academic.tex colour overrides
AMBER = RGBColor(0xB9, 0x77, 0x0E)
GREEN = RGBColor(0x1E, 0x84, 0x49)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x44, 0x44, 0x44)

NAME = "Debasish Chakroborti (Joy), Ph.D."
EMAIL = "debasish.chakroborti@usask.ca"
PHONE = "306 914 6212"
ADDR = "410 Stensrud Rd #118, Saskatoon, SK"
WEB = "https://joydeba.github.io/"
ORCID = "https://orcid.org/0000-0002-1597-8162"
GITHUB = "https://github.com/joydeba"
SCHOLAR = "https://scholar.google.com/citations?user=VdrsyT8AAAAJ"


def set_run(run, *, size=10, bold=False, italic=False, color=DARK, font="Georgia"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_horizontal_line(paragraph, color_hex="B9770E"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def setup_doc(margins=0.7):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(margins)
        section.bottom_margin = Inches(margins)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return doc


def header_block(doc, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(NAME)
    set_run(r, size=16, bold=True, color=DARK, font="Calibri")

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(subtitle)
        set_run(r2, size=11, bold=True, color=AMBER, font="Calibri")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    r = contact.add_run(f"{ADDR}  ·  {PHONE}  ·  {EMAIL}")
    set_run(r, size=9, color=MUTED, font="Calibri")

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(6)
    r = links.add_run(f"{WEB}  ·  {GITHUB}  ·  {ORCID}")
    set_run(r, size=9, color=GREEN, font="Calibri")
    add_horizontal_line(links)


def meta_line(doc, text):
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(text)
    set_run(r, size=9, italic=True, color=MUTED, font="Calibri")
    return meta


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run(r, size=11, bold=True, color=AMBER, font="Calibri")
    add_horizontal_line(p, "B9770E")
    return p


def sub_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, size=10, bold=True, color=GREEN, font="Calibri")
    return p


def body_para(doc, text, *, first_indent=False, space_after=6, size=10, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.2)
    if bold_lead:
        r0 = p.add_run(bold_lead)
        set_run(r0, size=size, bold=True, color=DARK, font="Georgia")
    r = p.add_run(text)
    set_run(r, size=size, color=DARK, font="Georgia")
    return p


def bullet(doc, text, *, bold_prefix=None, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        set_run(r0, size=size, bold=True, color=DARK, font="Georgia")
        r1 = p.add_run(text)
        set_run(r1, size=size, color=DARK, font="Georgia")
    else:
        r = p.add_run(text)
        set_run(r, size=size, color=DARK, font="Georgia")
    return p


def entry(doc, left, right_title, body=None):
    """Date-left style entry matching curve CV layout (simplified)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(3)
    r1 = p.add_run(f"{left}  ")
    set_run(r1, size=9, bold=True, color=GREEN, font="Calibri")
    r2 = p.add_run(right_title)
    set_run(r2, size=10, bold=True, color=DARK, font="Georgia")
    if body:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        p2.paragraph_format.left_indent = Inches(0.15)
        r = p2.add_run(body)
        set_run(r, size=9.5, color=MUTED, font="Georgia")


# ---------------------------------------------------------------------------
# 1. COVER LETTER
# ---------------------------------------------------------------------------

def write_cover_letter():
    doc = setup_doc(0.8)
    header_block(doc, "Cover Letter — Assistant Professor, Software Engineering & AI / Applied AI")
    meta_line(doc, "Requisition ID 2649  ·  School of Computer Science  ·  University of Guelph")

    body_para(doc, "August 2026", space_after=8)
    body_para(
        doc,
        "Dr. Minglun Gong\nDirector, School of Computer Science\nCollege of Computational Mathematical and Physical Sciences\nUniversity of Guelph, Guelph, ON  N1G 2W1\nfacultyjobs@socs.uoguelph.ca",
        space_after=8,
    )
    body_para(doc, "Dear Dr. Gong and Members of the Search Committee,", space_after=8)

    body_para(
        doc,
        "I am writing to apply for the tenure-track Assistant Professor position in Software Engineering & AI / "
        "Applied AI in the School of Computer Science at the University of Guelph (Requisition ID 2649). I hold a "
        "Ph.D. in Computer Science (Software Engineering) from the University of Saskatchewan, where I now serve as "
        "a Living Skies Postdoctoral Fellow with the NSERC CREATE program in Software Analytics Research (SOAR), "
        "contributing to an NSERC Alliance Advantage industry–academia project. I am concurrently an Instructor at "
        "Saskatchewan Polytechnic and have served as an AI Automation Consultant with Cloud Linux Software, Inc. "
        "My research and teaching sit squarely within the restricted scope of this search: I build AI-assisted "
        "methods and tools that let software teams maintain and evolve systems safely on the stable versions people "
        "actually use. I am a Canadian citizen and am available to start on January 1, 2027.",
        first_indent=True,
    )
    body_para(
        doc,
        "Scholarship. My research program in intelligent release engineering combines mining software repositories, "
        "machine learning (CNNs and LLMs, including agentic approaches), and program analysis. During my Ph.D. and "
        "postdoctoral work I designed and built Baczer, ReBack, BackSlice, and BackTrans — systems that identify, "
        "recommend, and adapt backporting changes across large open-source ecosystems. This work has produced 24 "
        "research items to date, with publications in Automated Software Engineering, Empirical Software Engineering, "
        "the Journal of Software: Evolution and Process, ACM PACMHCI (EICS), and at ICSE, ICPC, MSR, CASCON, and "
        "IEEE Big Data, alongside three best-poster awards. As detailed in my Research Statement, I propose an "
        "independent, fundable five-year program with three themes — intelligent release engineering and quality "
        "assurance, AI-assisted developer workflows and technical debt, and industry-connected applied AI for "
        "software maintenance — with a staged funding plan targeting the NSERC Discovery Grant and Discovery Launch "
        "Supplement, Mitacs Accelerate, and NSERC Alliance, building on my current Alliance Advantage and Cloud "
        "Linux collaborations.",
        first_indent=True,
    )
    body_para(
        doc,
        "Fit with Guelph. My agenda complements SoCS strengths in software engineering, AI, and data science "
        "without duplicating them: release engineering, backporting, and maintenance-focused applied AI are "
        "high-impact areas where I would bring an established, distinctive niche. My earlier applied-AI systems work "
        "at the Plant Phenotyping and Imaging Research Centre (P2IRC/GIFS) — building RISP/RISPts for reusable "
        "scientific workflows and InsCount for image-based phenotyping — aligns naturally with Guelph's leadership "
        "in agri-food and digital agriculture and creates concrete cross-college collaboration opportunities, "
        "including with CARE-AI. I am eager to help SoCS serve its growing cohorts of 1,200+ undergraduate and "
        "100+ graduate students.",
        first_indent=True,
    )
    body_para(
        doc,
        "Teaching. I have a demonstrated commitment to undergraduate and graduate teaching. I carried full course "
        "responsibility as Sessional Lecturer for two sections of CMPT 145 (Principles of Computer Science) at the "
        "University of Saskatchewan, guiding three teaching assistants; I currently teach and redesign courses "
        "across systems administration, cloud computing, DevOps, networking, and web development at Saskatchewan "
        "Polytechnic, including course outlines I authored that received formal program approval; "
        "and as a Lecturer at Khulna University I taught structured programming, software engineering, "
        "and project courses for four years. I have completed 250+ hours of formal instructional training (LIFT, "
        "GTI FTUT, NIO, CETL/IQAC). I have mentored an M.Sc. student to co-authored publications (ICSE 2023, "
        "CASCON 2025) and supervised three undergraduate summer researchers — matching the School's preference for "
        "supervising students on scholarship-oriented projects. I am ready to teach core computer science courses "
        "at all levels, offer graduate courses in software engineering and applied AI, supervise M.Sc. and Ph.D. "
        "students, and carry the typical three-course annual load.",
        first_indent=True,
    )
    body_para(
        doc,
        "Service and community engagement. I served as President of the Computer Science Graduate Council, Lead "
        "Organizer of the ICSAC/Research Fest symposium, and GSA representative on university committees; I am a PC "
        "member for MSR 2026 (Data/Challenge track) and the SOAR Symposium 2025, and a reviewer for JSS, JSME/SMR, "
        "ASE, FSE, ICPC, ICSE SEIS, CHI, SANER, ICSME, and CASCON. My outreach includes leading the Digitized "
        "Technology event introducing high-school students to computing, and multi-year community volunteering. I "
        "have completed Reconciliation Education and Four Seasons of Reconciliation training, and I am committed to "
        "the University's inclusive-excellence and Indigenization goals under Bi-Naagward | It Comes into View.",
        first_indent=True,
    )
    body_para(
        doc,
        "My integrated application includes this letter, my curriculum vitae, a Teaching Statement (with philosophy, "
        "interests, experiences, and course evaluations), a Research Statement (with a summary of achievements, a "
        "future research proposal, and a funding plan), two representative papers, and complete contact information "
        "for three references. I would welcome the opportunity to discuss how my scholarship, teaching, and service "
        "can strengthen the School of Computer Science. Thank you for your time and consideration.",
        first_indent=True,
    )
    body_para(doc, "Sincerely,", space_after=18)
    body_para(doc, "Debasish Chakroborti (Joy), Ph.D.", space_after=2)
    body_para(doc, f"{EMAIL}  ·  {PHONE}  ·  {WEB}", space_after=2)

    path = OUT / "2649_CoverLetter.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2. CURRICULUM VITAE
# ---------------------------------------------------------------------------

def write_cv():
    doc = setup_doc(0.65)
    header_block(doc, "Curriculum Vitae (Academic)")
    meta_line(doc, "Application: Assistant Professor, Software Engineering & AI / Applied AI · Requisition ID 2649 · University of Guelph · August 2026")

    section_heading(doc, "Research Interests")
    body_para(
        doc,
        "Software Engineering & Applied AI — AI-assisted software maintenance and evolution; automated backporting "
        "and release engineering; technical debt and quality in stable releases; mining software repositories; "
        "LLM/ML and agentic approaches for code analysis, adaptation, and recommendation; scientific workflow and "
        "data reusability (FAIR); industry-connected software automation. Goal: deliver issue-free software on the "
        "versions people actually use, while training students through research-informed teaching and mentoring.",
        size=9.5,
    )

    section_heading(doc, "Education")
    entry(
        doc,
        "2020–2024",
        "Ph.D., Computer Science / Software Engineering — University of Saskatchewan, Canada",
        "Thesis: Automated Backporting for Efficient Version Management in Software Repositories. "
        "Advisors: Prof. Chanchal K. Roy and Prof. Kevin A. Schneider. Course mark 89%. Awards include Carl McCrosky "
        "Innovation Scholarship (2023), GSA Mark Kroeker Exceptional Student Leadership Award (2023), Ph.D. Citizenship "
        "Award (2024), Education Graduate Award (Friends Club of Saskatoon, 2024), Faculty/Departmental Scholarship, "
        "best poster awards (SOAR 2021, SEMLA 2022, ICSAC 2022), and Graduate Travel Awards (Seattle, Melbourne, Montréal).",
    )
    entry(
        doc,
        "2017–2019",
        "M.Sc., Computer Science / Big Data — University of Saskatchewan, Canada",
        "Thesis: An Intermediate Data-driven Methodology for Scientific Workflow Management System to Support Reusability. "
        "Advisors: Prof. Chanchal K. Roy and Prof. Kevin A. Schneider. Course mark 89%. Saskatchewan Innovation and "
        "Opportunity Scholarship; University of Saskatchewan Faculty Stipend.",
    )
    entry(
        doc,
        "2008–2012",
        "B.Sc., Computer Science and Engineering — Khulna University, Bangladesh",
        "Thesis: Web Service Performance Enhancement for Portable Devices Modifying SOAP Security Principle "
        "(Supervisor: Dr. Masud Rahman). CGPA 3.73/4.00 (1st class, 3rd position). Innovation Award; Dean's Merit / "
        "Merit List Scholarships (2008–2012); Honorable Mention, AUST Inter-University Programming Contest (2009).",
    )

    section_heading(doc, "Academic and Professional Appointments")
    entry(
        doc,
        "2024–Present",
        "Living Skies Postdoctoral Fellow, Department of Computer Science, University of Saskatchewan",
        "NSERC CREATE Software Analytics Research (SOAR). Automated software backporting and AI-driven development "
        "tools; contributing researcher on an NSERC Alliance Advantage industry–academia project (Software Research "
        "Lab, SRLab). Mentor of graduate and undergraduate researchers.",
    )
    entry(
        doc,
        "2024–Present",
        "Instructor, School of Computing and Digital Innovation, Saskatchewan Polytechnic (Regina Campus)",
        "Computer Systems Technology and Cloud Computing & Blockchain programs; designed/rewrote course and lab "
        "materials for most courses taught. Selected faculty participant, China Henan (ZJTIE) Transnational "
        "Education Project (2026).",
    )
    entry(
        doc,
        "2025–2026",
        "Consultant (AI Automation Team), Cloud Linux Software, Inc., USA",
        "Research advisory on AI-based automated backporting tooling for production Linux maintenance.",
    )
    entry(
        doc,
        "2025",
        "Sessional Lecturer, Department of Computer Science, University of Saskatchewan",
        "CMPT 145.3 Principles of Computer Science (Spring/Summer 2025, sections MT1 and MT2; on-campus). Full course "
        "responsibility including deferred/supplemental examinations; guided 3 teaching assistants.",
    )
    entry(
        doc,
        "2024",
        "Software Engineer (Intern), Push Interactions, Saskatoon",
        "Web application enhancement (https://www.wilger.net).",
    )
    entry(
        doc,
        "2023",
        "Ph.D. Intern (Remote), Synesis IT GmbH, Germany",
        "API planning and big-data databases for Traffic Management Intelligent Transportation Systems.",
    )
    entry(
        doc,
        "2017–2024",
        "Teaching Assistant / Research Assistant, University of Saskatchewan",
        "TA (PSAC 40004) for CMPT 141, 145, 214, 270, 280, 370, and Database Management Systems. RA at P2IRC/GIFS: "
        "built RISP and RISPts (Hadoop/Spark scientific-workflow data reusability) and InsCount (UNet-based imaging). "
        "Built SE research tools Baczer, ReBack (CNN), BackTrans (LLM), and BackSlice.",
    )
    entry(
        doc,
        "2013–2017",
        "Lecturer, CSE Discipline, Khulna University, Bangladesh (2017–2024 on study leave)",
        "Lectures, labs, assignments, and examinations; undergraduate thesis/project guidance; developed Khulna "
        "University, ICCIT, and Convocation 2016 registration websites.",
    )
    entry(
        doc,
        "2014–2016",
        "Part-time Lecturer, North Western University; Instructor (Part-time), HSTTI, Khulna",
        "Computer Science courses (BA program); ICT training courses for high-school teachers.",
    )
    entry(
        doc,
        "2012–2013",
        "Software Engineer — Divine IT Limited; Nascenia Limited, Bangladesh",
        "ERP systems (Django/SQL); SnapKnot wedding-photography platform (Ruby on Rails / ASP.NET).",
    )

    section_heading(doc, "Publications")
    body_para(
        doc,
        "24 research items (2015–2026) including 6 journal articles, 10 conference papers, 4 manuscripts in "
        "preparation, and 4 book chapters/awarded posters. Venues include ICSE, ICPC, MSR, CASCON, IEEE Big Data, "
        "Automated Software Engineering, Empirical Software Engineering, JSME, and ACM PACMHCI (EICS).",
        size=9.5,
    )
    sub_heading(doc, "Journal Articles")
    for t in [
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2024). ReBack: Recommending backports in social coding "
        "environments. Automated Software Engineering, 31(1), 18. https://doi.org/10.1007/s10515-024-00416-1",
        "Chakroborti, D., Nath, S., Schneider, K., & Roy, C. (2022). Release conventions of open-source software: "
        "An exploratory study. Journal of Software: Evolution and Process, e2499. https://doi.org/10.1002/smr.2499",
        "Chakroborti, D., Roy, B., & Nath, S. S. (2021). Designing for recommending intermediate states in a scientific "
        "workflow management system. Proc. ACM Hum.-Comput. Interact., 5(EICS). https://doi.org/10.1145/3457145",
        "Herbold, S., Trautsch, A., Ledel, B., et al., including Chakroborti, D. (2021). A fine-grained data set and "
        "analysis of tangling in bug fixing commits. Empirical Software Engineering.",
        "Bairagi, A. K., Mondal, S., & Chakroborti, D. (2017). Securing Bangla text communication using image "
        "steganography with dynamic substitution in IoT environment. Khulna University Studies, 14.",
        "Chakroborti, D., Roy, C., & Schneider, K. (2024). Code propagation in stable version management: A systematic "
        "literature review. Manuscript under review.",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "Conference Proceedings")
    for t in [
        "Chakroborti, D. (2026). Backporting battles on the branches: Agentic vs human fixes. Submitted to ICSE.",
        "Tasnim, J., Chakroborti, D., Roy, C., & Schneider, K. (2025). An insight into the technical debt–fix "
        "trade-off in software backporting. CASCON 2025, Toronto, ON, Canada: IEEE.",
        "Chakroborti, D., Roy, C. K., & Schneider, K. A. (2024). A study of backporting code in open-source software "
        "for characterizing changesets. ICSE-Companion '24. https://doi.org/10.1145/3639478.3643079",
        "Chakroborti, D., Roy, C., & Schneider, K. (2024). BackSlice: Achieving accurate and essential propagation "
        "through backport slicing. Submitted.",
        "Tasnim, J., Chakroborti, D., Roy, C. K., & Schneider, K. A. (2023). How does quality deviate in stable "
        "releases by backporting? ICSE 2023, Melbourne, Australia.",
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2022). Backports: Change types, challenges and strategies. "
        "ICPC 2022. https://doi.org/10.1145/3524610.3527920",
        "Bhattacharjee, A., Nath, S. S., Zhou, S., Chakroborti, D., Roy, B., Roy, C. K., & Schneider, K. (2020). An "
        "exploratory study to find motives behind cross-platform forks from Software Heritage dataset. MSR 2020. "
        "https://doi.org/10.1145/3379597.3387512",
        "Chakroborti, D., Mondal, M., Roy, B., Roy, C. K., & Schneider, K. A. (2018). Optimized storing of workflow "
        "outputs through mining association rules. IEEE Big Data 2018. https://doi.org/10.1109/BigData.2018.8622351",
        "Chakroborti, D., & Nath, S. S. (2017). Web service performance enhancement for portable devices modifying "
        "SOAP security principle. ICCIT 2017. https://doi.org/10.1109/ICCITECHN.2017.8281799",
        "Bairagi, A. K., & Chakroborti, D. (2015). Trust based D2D communications for accessing services in Internet "
        "of Things. ICCIT 2015. https://doi.org/10.1109/ICCITechn.2015.7488041",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "In Preparation")
    for t in [
        "Chakroborti, D. An insight into backporting changes: An empirical study with pull-based development.",
        "Chakroborti, D. BranchBench and BranchBridge: Patch-level evaluation for CVE backport adaptation.",
        "Chakroborti, D. Beyond correctness: Technical debt and carbon footprint of LLM-agent versus human backports.",
        "Tasnim, J., Chakroborti, D., Roy, C. K., & Schneider, K. A. Impact of backporting on the quality of stable "
        "software releases: A comparative analysis.",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "Book Chapter and Awarded Posters")
    for t in [
        "Chakroborti, D., Roy, B., Mondal, A., Mostaeen, G., Roy, C. K., Schneider, K. A., & Deters, R. (2020). A data "
        "management scheme for micro-level modular computation-intensive programs in big data platforms. In Data "
        "Management and Analysis (Springer). https://doi.org/10.1007/978-3-030-32587-9_9",
        "BEST POSTER (3rd), SEMLA 2022 — Backporting for version management: Automated changesets integration in "
        "software repositories.",
        "BEST POSTER (2nd), ICSAC 2022 — Detecting and visualizing quality aspects in backporting pull requests "
        "(with mentee J. Tasnim).",
        "BEST POSTER, SOAR Symposium 2021 — Identifying and integrating changesets in backporting.",
    ]:
        bullet(doc, t, size=9)
    body_para(doc, f"Full list: {WEB}research  ·  Google Scholar: {SCHOLAR}", size=9)

    section_heading(doc, "Research Grants & Funding")
    bullet(doc, "NSERC Alliance Advantage Grant — contributing Postdoctoral Fellow on an industry–academia project (ongoing).", size=9)
    bullet(doc, "NSERC CREATE SOAR — Living Skies Postdoctoral Fellowship, University of Saskatchewan (2024–present).", size=9)
    bullet(doc, "Cloud Linux Software Consultancy Fund — AI Automation Consultant (AI-based automated backporting).", size=9)
    bullet(doc, "Saskatchewan Polytechnic Professional Development Fund — conference presentation support (CASCON 2025, Toronto).", size=9)
    bullet(doc, "Graduate Travel Awards — international conference participation (Seattle 2018, Melbourne 2023, Montréal 2022).", size=9)
    bullet(doc, "Funding application support — assisted with drafting a portion of a research funding application (M.Sc. mentee project).", size=9)

    section_heading(doc, "Research Tools & Artifacts")
    bullet(doc, "Baczer — backport changeset identification and analysis at ecosystem scale.", bold_prefix="", size=9)
    bullet(doc, "ReBack — CNN-based recommendation of backports in social coding environments (ASE journal, 2024).", size=9)
    bullet(doc, "BackSlice — program-slicing approach for accurate, essential backport propagation.", size=9)
    bullet(doc, "BackTrans — LLM-based adaptation/translation of patches across versions and branches.", size=9)
    bullet(doc, "RISP / RISPts — intermediate-data reusability in scientific workflow management systems (Hadoop/Spark, FAIR).", size=9)
    bullet(doc, "InsCount — UNet-based image analysis tool for plant phenotyping (P2IRC/GIFS).", size=9)

    section_heading(doc, "Talks, Posters & Conference Visits")
    bullet(doc, "CASCON 2025 (Toronto) — talk: technical debt–fix trade-off in software backporting.", size=9)
    bullet(doc, "ICSE 2023 (Melbourne) — poster: quality deviation in stable releases by backporting.", size=9)
    bullet(doc, "SEMLA 2022 (Montréal) — Best Poster (3rd): Backporting for Version Management.", size=9)
    bullet(doc, "ICPC 2022 (with ICSE, Pittsburgh) — talk: Backports: Change Types, Challenges and Strategies.", size=9)
    bullet(doc, "IEEE Big Data 2018 (Seattle) — talk: Optimized Storing of Workflow Outputs through Mining Association Rules.", size=9)
    bullet(doc, "ICCIT 2017 / 2015 (Dhaka) — talks on web-service performance and trust-based D2D communications.", size=9)
    bullet(doc, "SOAR Symposium 2021 (Saskatoon) — Best Poster; CSGC Research Fest / ICSAC distinguished speaker recognition.", size=9)

    section_heading(doc, "Student Supervision & Mentoring")
    bullet(doc, "Jarin Tasnim (M.Sc., USask) — software backporting and quality analysis; co-authored ICSE 2023 and CASCON 2025 papers and an awarded ICSAC poster.", bold_prefix="M.Sc. mentoring — ", size=9)
    bullet(doc, "Supervised 3 undergraduate summer research students on scholarship-oriented projects.", bold_prefix="Undergraduate research — ", size=9)
    bullet(doc, "Guided 3 TAs as Sessional Lecturer for CMPT 145 (USask, Spring/Summer 2025).", bold_prefix="Teaching assistants — ", size=9)

    section_heading(doc, "Teaching Experience")
    entry(
        doc,
        "2024–Present",
        "Instructor, Saskatchewan Polytechnic — Computer Systems Technology; Cloud Computing & Blockchain",
        "COOS 190 Server Administration; COOS 294 Cloud Infrastructure Administration; COOS 295 Systems Administration 2; "
        "COET 295 Emerging Technologies; CNET 184 Data Communications and Networking; CWEB 280 Internet Programming and "
        "Web Applications 2; CMPT 145 Principles of Computer Science; CCMP 600–606 (cloud fundamentals, blockchain, "
        "provisioning/security, smart contracts, orchestration, cloud data management, integrated services); DEVP 600 "
        "DevOps; PROJ 611 industry cloud-adoption capstone. Designed/rewrote course and lab materials for most "
        "offerings; authored course outlines for CNET 184 and CWEB 280 that received program approval.",
    )
    entry(
        doc,
        "2025",
        "Sessional Lecturer, University of Saskatchewan — CMPT 145.3 (two sections, ~38 students/section)",
        "Full course responsibility; SLEQ end-of-course report available (response rate below release threshold for "
        "closed-ended scores; qualitative comments on file).",
    )
    entry(
        doc,
        "2017–2024",
        "Teaching Assistant, University of Saskatchewan",
        "CMPT 141, 145, 214, 270, 280, 370, Database Management Systems — tutorials, labs, marking, help desk, invigilation.",
    )
    entry(
        doc,
        "2013–2017",
        "Lecturer, Khulna University — CSE Discipline",
        "Structured Programming; Software Development Project; Advanced Programming Laboratory; Software Engineering "
        "and Information System (lectures, labs, assignments, examinations).",
    )

    section_heading(doc, "Academic Service & Peer Review")
    sub_heading(doc, "Program & Organizing Committees")
    bullet(doc, "PC Member, Mining Software Repositories (MSR 2026) — Data/Challenge track.", size=9)
    bullet(doc, "PC Member, SOAR Symposium 2025 (University of Saskatchewan).", size=9)
    bullet(doc, "Lead Organizer, 7th Symposium on Innovations in Computer Science and Applied Computing (ICSAC / Research Fest).", size=9)
    bullet(doc, "Judge: Research Fest 2025; USask Images of Research (2024); GSA Elevator Pitch (2024); GSA Research Conference (2023).", size=9)
    sub_heading(doc, "Journal & Conference Reviewing")
    bullet(doc, "Journals: Journal of Software: Evolution and Process (2023–2025); Journal of Systems and Software (2020, 2022, 2025, 2026).", size=9)
    bullet(doc, "Conferences: ASE 2026/2024/2020; FSE 2026; ICPC 2026; ICSE 2026 SEIS; CHI 2025; SANER 2025/2023; CASCON 2024; ICSME 2023; ISEC 2022; ICSE 2020 Tool Demo; IWSC 2020; ICSA 2020.", size=9)

    section_heading(doc, "Leadership, Outreach & Community Engagement")
    bullet(doc, "President, Computer Science Graduate Council (CSGC), USask (2022–2023); earlier VP–Admin and VP–Social.", size=9)
    bullet(doc, "GSA Representative — Associate Dean (Students) Search Committee; Arts & Science Planning Advisory Committee.", size=9)
    bullet(doc, "Volunteer Lead, Digitized Technology — USask computing outreach event for high-school students.", size=9)
    bullet(doc, "Volunteer Lead, Saskatchewan Student Partnership Program (SSPP); COVID-19 student support coordinator (2020–2022).", size=9)
    bullet(doc, "Community volunteering: Saskatoon Folkfest (2021–2025), Global Village Festival, cultural associations; club leadership (Huskies Cricket Club President, 2023–2024).", size=9)

    section_heading(doc, "Awards & Scholarships (selected)")
    bullet(doc, "Ph.D. Citizenship Award, University of Saskatchewan (2024); Education Graduate Award, Friends Club of Saskatoon (2024).", size=9)
    bullet(doc, "Carl McCrosky Innovation Scholarship (2023); GSA Mark Kroeker Exceptional Student Leadership Award (2023).", size=9)
    bullet(doc, "Best Poster Awards: SOAR 2021; SEMLA 2022 (3rd); ICSAC 2022 (2nd, with mentee).", size=9)
    bullet(doc, "Saskatchewan Innovation and Opportunity Scholarship; Departmental Scholarship; Ivan and Margaret Toutloff Award.", size=9)
    bullet(doc, "Innovation Award and Dean's Merit Scholarships, Khulna University (2008–2012).", size=9)

    section_heading(doc, "Certifications & Professional Development (selected)")
    bullet(doc, "LIFT Certificate, Saskatchewan Polytechnic — 90 hours, completed May 2026 (interactive lecturing, formative assessment, Indigenization & reconciliation, curriculum, learning technologies) with instructional ePortfolio.", size=9)
    bullet(doc, "New Instructor Orientation, Saskatchewan Polytechnic — 18 hours (2025).", size=9)
    bullet(doc, "GTI Foundation Training for University Teachers — 150 hours (2014); CETL/IQAC pedagogy modules (2014–2016).", size=9)
    bullet(doc, "Reconciliation Education; Four Seasons of Reconciliation; TCPS 2: CORE 2022 (research ethics).", size=9)
    bullet(doc, "AWS Academy Educator; AWS Skill Builder badges (Cloud, Blockchain, DevOps, ECS/EKS, Backup, etc.).", size=9)
    bullet(doc, "Coursera Machine Learning and Deep Learning Specializations; ACM and IEEE memberships.", size=9)

    section_heading(doc, "Technical Skills")
    body_para(
        doc,
        "Languages: Python, Java, C/C++, C#, Ruby, JavaScript, R, PHP, MATLAB, SQL, LaTeX. Frameworks and platforms: "
        "Django, Rails, Flask, ASP.NET, Spark, Hadoop, TensorFlow, PyTorch, AWS, Docker/Kubernetes, Git/GitHub CI. "
        "Methods: mining software repositories, empirical SE, ML/LLM and agentic pipelines for code, program analysis "
        "and slicing, statistical analysis. Human languages: English, Bengali.",
        size=9,
    )

    section_heading(doc, "References")
    body_para(doc, "Names and complete contact information for three references are provided in the accompanying References document.", size=9.5)

    path = OUT / "2649_CV.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 3. TEACHING STATEMENT
# ---------------------------------------------------------------------------

def write_teaching_statement():
    doc = setup_doc(0.8)
    header_block(doc, "Teaching Statement")
    meta_line(
        doc,
        "Application: Assistant Professor, Software Engineering & AI / Applied AI\n"
        "School of Computer Science, University of Guelph · Requisition ID 2649 · August 2026\n"
        "(Teaching philosophy · interests · experiences · course evaluations)"
    )

    section_heading(doc, "1. Teaching Philosophy")
    body_para(
        doc,
        "I teach so that students can design, build, test, and reason about software systems — including systems "
        "with AI components — with rigour and responsibility. Three principles guide my practice. First, authentic "
        "practice: students learn software engineering by doing it the way professionals do, so my courses embed "
        "version control, code review, automated testing, CI/CD, and staged team projects from the first weeks. In "
        "my Saskatchewan Polytechnic DevOps and cloud courses, every lab runs against real cloud infrastructure and "
        "every project is submitted through a pipeline the students configure themselves. Second, transparency and "
        "feedback: I use outcomes-based course design, published rubrics, worked examples, and frequent low-stakes "
        "formative assessment so learners always know what success looks like and how to close the gap. My LIFT "
        "training (90 hours) in interactive lecturing and formative assessment directly shaped how I run active "
        "classrooms — think-pair-share, live coding with deliberate mistakes, peer code walkthroughs, design "
        "reviews, and reflective project postmortems that feed back into the next iteration. In software engineering "
        "and design courses I foreground the full lifecycle — requirements, architecture, implementation, testing "
        "and quality assurance, maintenance and evolution — alongside problem framing, constraints, iterative "
        "prototyping, and risk assessment in team settings. Third, research-informed teaching: I bring my own studies of "
        "real open-source ecosystems into class, so students analyze genuine pull requests, releases, and failures "
        "rather than toy examples, and see how classroom concepts connect to the systems the world depends on.",
        first_indent=True,
    )

    section_heading(doc, "2. Teaching Interests and Readiness at Guelph")
    body_para(
        doc,
        "The posting asks for the ability to teach core computer science undergraduate courses at all levels and "
        "graduate courses in the candidate's scholarship area. My interests and record cover that full span.",
        first_indent=True,
    )
    bullet(doc, "introductory programming and principles of computer science (I have taught this level with full course responsibility);", bold_prefix="Core foundations — ", size=9.5)
    bullet(doc, "data structures, systems programming, databases, and networking (TA and instructor experience across CMPT 141–370 and CNET/COOS offerings);", bold_prefix="Intermediate core — ", size=9.5)
    bullet(doc, "software engineering, software design and architecture, requirements analysis, system analysis, testing and quality assurance, software project management, and team-based software projects/capstones — the area of my research and four years of lecturing at Khulna University;", bold_prefix="Software engineering — ", size=9.5)
    bullet(doc, "cloud computing, DevOps, and web development, where I currently teach and have rewritten curricula (COOS/CCMP/DEVP/CWEB series, AWS Academy Educator);", bold_prefix="Applied systems — ", size=9.5)
    bullet(doc, "mining software repositories and software analytics; AI-assisted software maintenance and release engineering; engineering AI-enabled systems (LLM/agent-based development workflows, their quality, technical debt, and sustainability).", bold_prefix="Graduate topics — ", size=9.5)
    body_para(
        doc,
        "At Guelph this maps naturally onto core programming and data-structures courses, software engineering and "
        "system analysis/design offerings, team project courses, and new graduate seminars in software analytics and "
        "applied AI for software maintenance that would complement SoCS's existing graduate programs. I am "
        "comfortable carrying the typical three-course annual load across this range, and I keep all materials "
        "current with modern tooling and with a clear, principled policy on student use of AI assistants: students "
        "learn to use them critically — and to verify, test, and take responsibility for what they produce.",
        first_indent=True,
    )

    section_heading(doc, "3. Teaching Experience")
    body_para(
        doc,
        "Full course responsibility. As Sessional Lecturer at the University of Saskatchewan (Spring/Summer 2025) I "
        "taught two on-campus sections of CMPT 145.3 Principles of Computer Science, with complete responsibility "
        "for lectures, assessments, deferred and supplemental examinations, and re-reads, while supervising three "
        "teaching assistants. Coordinating parallel sections taught me to keep large first-year cohorts consistent, "
        "fair, and engaged — directly relevant to SoCS's 1,200+ undergraduate enrolment.",
        first_indent=True,
    )
    body_para(
        doc,
        "Polytechnic instruction and curriculum design. As Instructor at Saskatchewan Polytechnic (2024–present) I "
        "teach across two programs — Computer Systems Technology (COOS 190 Server Administration, COOS 294 Cloud "
        "Infrastructure Administration, COOS 295 Systems Administration 2, COET 295 Emerging Technologies, CNET 184 "
        "Data Communications, CWEB 280 Internet Programming 2, CMPT 145) and Cloud Computing & Blockchain "
        "(CCMP 600–606, DEVP 600 DevOps, PROJ 611 industry capstone). I treat course development as scholarly work: I designed or rewrote the course "
        "and lab materials for most of these offerings, and I authored course outlines for CNET 184 and CWEB 280 "
        "that received formal program approval — direct experience with outcomes-based curriculum processes. I was "
        "also selected as a faculty participant in the China Henan (ZJTIE) transnational education project (2026), "
        "teaching and running labs abroad.",
        first_indent=True,
    )
    body_para(
        doc,
        "University lecturing and assistantship. As Lecturer in CSE at Khulna University (2013–2017) I taught "
        "Structured Programming, Software Development Project, Advanced Programming Laboratory, and Software "
        "Engineering and Information System, including undergraduate project guidance. As TA at the University of "
        "Saskatchewan (2017–2024) I supported CMPT 141, 145, 214, 270, 280, 370, and Database Management Systems — "
        "tutorials, labs, marking, and help-desk support across the whole undergraduate core.",
        first_indent=True,
    )

    section_heading(doc, "4. Mentorship and Scholarship-Oriented Supervision")
    body_para(
        doc,
        "This search prefers candidates with experience supervising students or staff on scholarship-oriented "
        "projects. I mentored M.Sc. student Jarin Tasnim (University of Saskatchewan) from problem formulation to "
        "publication: our co-authored work on quality deviation in stable releases appeared at ICSE 2023, our study "
        "of the technical debt–fix trade-off in backporting at CASCON 2025, and her ICSAC 2022 poster won a "
        "best-poster award. I have also supervised three undergraduate summer research students and guided three "
        "teaching assistants. My mentoring model uses weekly milestone meetings, reproducible research artifacts "
        "from day one, early writing, and conference presentation as an explicit goal. At Guelph I will supervise "
        "M.Sc. and Ph.D. students in SoCS's graduate programs and create funded undergraduate research roles "
        "(e.g., NSERC USRA) within my lab's projects.",
        first_indent=True,
    )

    section_heading(doc, "5. Course Evaluations and Evidence of Effectiveness")
    body_para(
        doc,
        "For CMPT 145 (USask, Spring/Summer 2025), the official SLEQ end-of-course report is available; because only "
        "3 of 38 students responded (7.9%, below the release threshold), closed-ended scores were not released and "
        "the report contains qualitative comments only — I include it for completeness and transparency: "
        "https://joydeba.github.io/teaching-materials/sleq-cmpt145-2025.pdf. Broader evidence of effectiveness "
        "includes my full teaching dossier (philosophy, course portfolio, materials, EDIA, and mentorship evidence: "
        "https://joydeba.github.io/teaching-materials/teaching-dossier.pdf), the LIFT Statement of Achievement and "
        "instructional ePortfolio, repeat teaching assignments and curriculum-design responsibility at Saskatchewan "
        "Polytechnic, and selection for the ZJTIE international teaching project. I have completed 250+ hours of "
        "formal instructional training: LIFT (90 h, 2026), New Instructor Orientation (18 h, 2025), GTI Foundation "
        "Training for University Teachers (150 h, 2014), and CETL/IQAC pedagogy modules (2014–2016).",
        first_indent=True,
    )

    section_heading(doc, "6. Inclusive Teaching and Reconciliation")
    body_para(
        doc,
        "Inclusive teaching is not optional. I completed Reconciliation Education, Four Seasons of Reconciliation, "
        "and LIFT's Indigenization & Reconciliation module, and I apply them concretely: transparent expectations "
        "and rubrics, multimodal materials, flexible low-stakes assessment structures, and classroom norms that make "
        "asking for help safe. Having taught in three countries and supported international cohorts (including "
        "coordinating student support during COVID-19), I design for students who arrive with very different prior "
        "preparation. I am committed to the University of Guelph's inclusive-excellence goals and to contributing "
        "to reconciliation through action, guided by Bi-Naagward | It Comes into View. My aim is a classroom that "
        "is rigorous, welcoming, and practice-grounded — preparing SoCS's growing cohorts for careers and research "
        "in software engineering and AI.",
        first_indent=True,
    )

    path = OUT / "2649_TeachingStatement.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4. RESEARCH (SCHOLARSHIP) STATEMENT
# ---------------------------------------------------------------------------

def write_research_statement():
    doc = setup_doc(0.8)
    header_block(doc, "Research Statement (Scholarship Statement)")
    meta_line(
        doc,
        "Application: Assistant Professor, Software Engineering & AI / Applied AI\n"
        "School of Computer Science, University of Guelph · Requisition ID 2649 · August 2026\n"
        "(Summary of research achievements · proposal for future research · funding plan)"
    )

    section_heading(doc, "1. Vision and Fit")
    body_para(
        doc,
        "I work at the intersection of software engineering and applied AI. My goal is practical and ambitious: "
        "deliver issue-free software on the versions people actually use. Most of the world's critical software runs "
        "on stable releases and long-lived branches — not on the latest development head — and that is where fixes, "
        "security patches, and quality problems are hardest to manage. Backporting a fix to a stable branch is "
        "high-stakes, expert work that today is largely manual. I design intelligent, data-driven methods and tools "
        "that help developers anticipate risk, propagate safe changes, manage technical debt, and evolve software "
        "efficiently — and I evaluate them empirically on large open-source ecosystems and in industry. This agenda "
        "sits squarely within the restricted scope of this search (Software Engineering and AI), complements SoCS's "
        "existing strengths with a distinctive niche in release engineering and AI-assisted maintenance, and is "
        "designed from the outset as an independent, externally fundable faculty program.",
        first_indent=True,
    )

    section_heading(doc, "2. Research Achievements and Impact")
    body_para(
        doc,
        "A. A complete pipeline for AI-assisted backporting. In my Ph.D. at the University of Saskatchewan "
        "(advisors Prof. Chanchal K. Roy and Prof. Kevin A. Schneider), I built the first end-to-end research "
        "pipeline for automated backporting: Baczer characterizes backporting changesets at ecosystem scale "
        "(ICPC 2022; ICSE-Companion 2024); ReBack uses convolutional neural networks over social-coding signals to "
        "recommend which pull requests should be backported (Automated Software Engineering, 2024); BackSlice uses "
        "program slicing to propagate only the accurate, essential parts of a change; and BackTrans uses large "
        "language models to adapt patches across diverging branches. Together these systems attack every stage of a "
        "maintenance task that affects nearly every long-lived software product. My empirical work also established "
        "how release conventions actually behave in open source (JSME, 2022) and contributed to a widely used "
        "fine-grained dataset of tangled bug-fixing commits (Empirical Software Engineering, 2021).",
        first_indent=True,
    )
    body_para(
        doc,
        "B. Quality, technical debt, and the human–AI comparison. With my M.Sc. mentee I quantified how quality "
        "deviates in stable releases through backporting (ICSE 2023) and characterized the technical debt–fix "
        "trade-off maintainers accept when they backport (CASCON 2025). My current work extends this to the agentic "
        "era: a submitted ICSE study compares LLM-agent versus human backports head-to-head, and in-preparation work "
        "(BranchBench/BranchBridge) contributes patch-level benchmarks for CVE backport adaptation and measures the "
        "technical debt and carbon footprint of agentic fixes. This positions me at the front of a question every "
        "software organization is now asking: where can AI maintenance be trusted, and at what long-term cost?",
        first_indent=True,
    )
    body_para(
        doc,
        "C. Applied AI systems beyond one niche. At the Plant Phenotyping and Imaging Research Centre (P2IRC/GIFS) "
        "I built RISP and RISPts, which make intermediate data in scientific workflows findable and reusable on "
        "Hadoop/Spark (IEEE Big Data 2018; ACM PACMHCI/EICS 2021; Springer chapter 2020), and InsCount, a UNet-based "
        "imaging tool for phenotyping. This demonstrates that I deliver usable research software for interdisciplinary "
        "partners — experience directly relevant to Guelph's agri-food and digital-agriculture research environment.",
        first_indent=True,
    )
    body_para(
        doc,
        "D. Industry translation and training. As Living Skies Postdoctoral Fellow (NSERC CREATE SOAR) I contribute "
        "to an NSERC Alliance Advantage industry–academia project, and as AI Automation Consultant with Cloud Linux "
        "Software, Inc. I have applied my backporting research to production Linux maintenance. In total my record "
        "comprises 24 research items with three best-poster awards, and co-authored publications with a mentored "
        "M.Sc. student and undergraduate researchers — evidence that I already run scholarship-oriented student "
        "projects and can carry collaborations from idea to deployed practice.",
        first_indent=True,
    )

    section_heading(doc, "3. Proposed Research Program (Years 1–5)")
    body_para(
        doc,
        "My lab — the Intelligent Release Engineering (IRE) Lab — will pursue three interlocking themes. Each theme "
        "has near-term projects that produce publishable results within 12–18 months and longer-horizon goals that "
        "sustain a five-year program.",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 1 — Intelligent release engineering and quality assurance. Near term: risk models that forecast, at "
        "pull-request time, whether a change will need backporting and what its quality impact on stable releases "
        "will be, built on my Baczer/ReBack foundations; public benchmarks (BranchBench) for patch-level backport "
        "adaptation, including CVE fixes. Longer term: a release-engineering assistant that recommends testing and "
        "remediation pathways per branch, evaluated with industrial partners. Deliverables: open datasets and "
        "benchmarks, risk models, developer-facing tools, and publications in top SE venues (ICSE, FSE, ASE, EMSE, TSE).",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 2 — Trustworthy AI-assisted maintenance and technical debt. Near term: extend my agentic-versus-human "
        "backporting studies into a general framework for auditing LLM/agent maintenance — correctness, maintainability, "
        "technical debt, reproducibility, and energy/carbon cost. Longer term: hybrid human–AI workflows in which "
        "agents propose and humans govern, with empirically derived guardrails; guidelines that practitioners and "
        "educators can apply directly. A complementary strand studies how design decisions, process choices, and "
        "collaboration patterns influence maintainability and delivery outcomes in team-based development — in both "
        "professional and student teams — which links this theme directly to my teaching of software engineering "
        "and engineering AI-enabled systems.",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 3 — Industry-connected applied AI for software maintenance and digital agriculture. Near term: "
        "Mitacs-supported projects with existing partners (building on Cloud Linux and Alliance Advantage ties) on "
        "legacy modernization and secure maintenance automation. Longer term: leverage my P2IRC/GIFS experience to "
        "partner with Guelph's agri-food and digital-agriculture community — where research software, long-lived "
        "field systems, and scientific workflows raise exactly the maintenance and data-reusability problems I "
        "study — and engage with CARE-AI on responsible-AI dimensions of automated maintenance. Deliverables: "
        "funded partnerships, deployed tools, and translational publications.",
        first_indent=True,
    )

    section_heading(doc, "4. Funding Plan (Tri-Council and Related)")
    body_para(
        doc,
        "The position expects an internationally recognized program funded by NSERC or similarly prestigious "
        "agencies. My plan is staged and realistic, and builds on funding relationships I already hold:",
        first_indent=True,
    )
    bullet(
        doc,
        "NSERC Discovery Grant plus Discovery Launch Supplement (early-career) as the program's foundation, "
        "submitted in my first fall; institutional start-up to equip the lab and recruit the first cohort; a Mitacs "
        "Accelerate internship with an existing industry partner to generate early applied results.",
        bold_prefix="Year 1–2: ",
        size=9.5,
    )
    bullet(
        doc,
        "NSERC Alliance proposal with industry partners (extending my Cloud Linux and Alliance Advantage "
        "collaborations, and new Ontario partners) on maintenance automation; CFI JELF consideration for compute "
        "infrastructure if needed; additional Mitacs placements; NSERC USRA students each summer.",
        bold_prefix="Year 2–4: ",
        size=9.5,
    )
    bullet(
        doc,
        "Discovery renewal at increased scale; larger partnered grants (Alliance, OCI/industry consortia) and "
        "international collaboration; sustained pipeline of 2–3 Ph.D. and 2–3 M.Sc. students plus undergraduates.",
        bold_prefix="Year 4–5: ",
        size=9.5,
    )
    bullet(
        doc,
        "a Discovery/Alliance-class budget typically supports 1–2 Ph.D. and 1–2 M.Sc. students, computing, open "
        "artifact curation, and dissemination. My research is compute-moderate (repository mining and fine-tuning "
        "rather than frontier-model training), which keeps budgets credible and results reproducible.",
        bold_prefix="Scale and feasibility: ",
        size=9.5,
    )
    body_para(
        doc,
        "I have direct experience inside this funding ecosystem: I contribute to an NSERC Alliance Advantage project "
        "as a postdoctoral researcher, hold an NSERC CREATE-affiliated fellowship, have industry consultancy funding, "
        "and have assisted in drafting grant application materials. I will also pursue smaller instruments "
        "opportunistically (SSHRC-relevant collaborations on human aspects of AI-assisted work, conference and "
        "travel programs, and institutional seed funds).",
        first_indent=True,
    )

    section_heading(doc, "5. Student Training and Open Science")
    body_para(
        doc,
        "Students are the core output of the program. Every thesis project is scoped to produce an open artifact "
        "(dataset, benchmark, or tool) plus a publication, following the mentoring model that already produced "
        "co-authored ICSE and CASCON papers with my M.Sc. mentee and an awarded student poster. I will supervise "
        "M.Sc. and Ph.D. students in SoCS's graduate programs, involve undergraduates through USRA and project "
        "courses, and integrate research modules into my software engineering and applied-AI teaching — consistent "
        "with the position's 40% scholarship / 40% teaching / 20% service balance. All tools and datasets are "
        "released openly (as with my existing artifacts), and empirical studies follow TCPS 2 ethics practice.",
        first_indent=True,
    )

    section_heading(doc, "6. Collaboration and Broader Impact")
    body_para(
        doc,
        "Within SoCS I see natural collaborations in software engineering, AI/ML, data science, and cybersecurity "
        "(secure patch propagation); across the College of Computational Mathematical and Physical Sciences, in "
        "statistics and mathematical modelling of software quality; and across campus, in agri-food informatics and "
        "responsible AI (CARE-AI). Beyond academia, safer stable-release maintenance means fewer outages and "
        "unpatched vulnerabilities in the software people depend on — infrastructure, health, and agriculture "
        "included. That is my discipline's version of the University of Guelph's purpose: To Improve Life.",
        first_indent=True,
    )

    path = OUT / "2649_ResearchStatement.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 5. REFERENCES
# ---------------------------------------------------------------------------

def write_references():
    doc = setup_doc(0.85)
    header_block(doc, "References")
    meta_line(
        doc,
        "Assistant Professor Application · Requisition ID 2649 · School of Computer Science, University of Guelph\n"
        "Names and complete contact information for three references, as required by the posting."
    )

    def referee(n, name, title, inst, email, phone=None, rel=None):
        section_heading(doc, n)
        body_para(doc, name, space_after=2)
        body_para(doc, title, space_after=2)
        body_para(doc, inst, space_after=2)
        line = f"Email: {email}"
        if phone:
            line += f"    Phone: {phone}"
        body_para(doc, line, space_after=2)
        if rel:
            body_para(doc, f"Relationship: {rel}", space_after=8)

    referee(
        "Referee 1",
        "Prof. Dr. Chanchal K. Roy",
        "Professor of Software Engineering, Department of Computer Science",
        "University of Saskatchewan, 176 Thorvaldson Building, 110 Science Place, Saskatoon, SK S7N 5C9, Canada",
        "chanchal.roy@usask.ca",
        phone="+1 (306) 715-0600",
        rel="Ph.D. and M.Sc. co-advisor; current postdoctoral supervisor (SRLab)",
    )
    referee(
        "Referee 2",
        "Prof. Dr. Kevin A. Schneider",
        "Professor, Department of Computer Science",
        "University of Saskatchewan, 176 Thorvaldson Building, 110 Science Place, Saskatoon, SK S7N 5C9, Canada",
        "kevin.schneider@usask.ca",
        phone="+1 (306) 966-4891",
        rel="Ph.D. and M.Sc. co-advisor",
    )
    referee(
        "Referee 3",
        "Dr. Banani Roy",
        "Associate Professor, Department of Computer Science",
        "University of Saskatchewan, 176 Thorvaldson Building, 110 Science Place, Saskatoon, SK S7N 5C9, Canada",
        "banani.roy@usask.ca",
        phone="+1 (306) 850-5630",
        rel="M.Sc. research mentor (scientific workflow management / P2IRC projects; co-authored publications)",
    )

    path = OUT / "2649_References.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 6. REPRESENTATIVE PAPERS (COVER SHEET)
# ---------------------------------------------------------------------------

def write_papers_included():
    doc = setup_doc(0.85)
    header_block(doc, "Two Representative Papers")
    meta_line(
        doc,
        "Assistant Professor Application · Requisition ID 2649 · University of Guelph\n"
        "The two full papers follow this cover sheet in the integrated application PDF."
    )

    section_heading(doc, "Why I Am Selecting These Two Papers")
    body_para(
        doc,
        "I am selecting these two papers because they revive backporting research in pull-based development. "
        "Previously, backporting research was concentrated mainly on cross-system development and the Linux kernel. "
        "These two works show that backporting is becoming common in pull-based development—and that we need "
        "automation to support it. Agentic pull-request development is now creating even more versions and branches, "
        "which increases backporting demand; I am studying that scenario in my submitted ICSE paper "
        "Backporting Battles on the Branches: Agentic vs Human Fixes. I would welcome the opportunity to "
        "discuss this work further in an interview.",
        first_indent=True,
    )

    section_heading(doc, "Paper 1 — Journal (applied AI for software engineering)")
    body_para(
        doc,
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2024). ReBack: Recommending backports in social coding "
        "environments. Automated Software Engineering, 31(1), 18. https://doi.org/10.1007/s10515-024-00416-1",
        space_after=6,
    )
    body_para(
        doc,
        "Summary: ReBack learns from social-coding signals (pull requests, reviews, discussion, and change "
        "characteristics) using convolutional neural networks to recommend which changes should be backported to "
        "stable branches, evaluated on large open-source ecosystems.",
        first_indent=True,
    )
    body_para(
        doc,
        "Why representative: Peer-reviewed journal evidence of machine learning applied to a core, high-stakes "
        "software maintenance task — the exact combination of Software Engineering & AI / Applied AI this search "
        "targets — and the centrepiece of the research program proposed in my Research Statement. "
        "(PDF in folder: s10515-024-00416-1.pdf)",
        first_indent=True,
    )

    section_heading(doc, "Paper 2 — Conference (empirical foundation)")
    body_para(
        doc,
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2022). Backports: Change types, challenges and strategies. "
        "In 30th IEEE/ACM International Conference on Program Comprehension (ICPC 2022), Pittsburgh, PA, USA. "
        "https://doi.org/10.1145/3524610.3527920",
        space_after=6,
    )
    body_para(
        doc,
        "Summary: A large-scale empirical study of backporting practice in open-source repositories — what kinds of "
        "changes get backported, what makes them hard, and what strategies maintainers use — establishing the "
        "evidence base for automating backport identification, propagation, and adaptation.",
        first_indent=True,
    )
    body_para(
        doc,
        "Why representative: First-author empirical work at a premier program-comprehension venue (co-located with "
        "ICSE) that grounds the AI-assisted maintenance agenda in how developers actually work. "
        "(PDF in folder: 3524610.3527920.pdf)",
        first_indent=True,
    )

    section_heading(doc, "Assembly note")
    body_para(
        doc,
        "In the single integrated PDF, place this sheet after the Research Statement, followed immediately by the "
        "two paper PDFs (s10515-024-00416-1.pdf, then 3524610.3527920.pdf), and end with the References document. "
        "Optional swap for Paper 2 if a mentoring-focused sample is preferred: Tasnim, J., Chakroborti, D., Roy, C., "
        "& Schneider, K. (2025), An insight into the technical debt–fix trade-off in software backporting, CASCON 2025.",
        space_after=6,
    )

    path = OUT / "2649_PapersIncluded.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [
        write_cover_letter(),
        write_cv(),
        write_teaching_statement(),
        write_research_statement(),
        write_references(),
        write_papers_included(),
    ]
    for p in paths:
        print(p)


if __name__ == "__main__":
    main()
