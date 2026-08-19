#!/usr/bin/env python3
"""Generate Haverford College CS TT Assistant Professor application DOCX
files, styled to match Debasish Chakroborti academic CV (amber/green accents).

Position: Tenure-Track Assistant Professor of Computer Science
Department of Computer Science, Haverford College. Start Fall 2027.
Deadline: 11:59pm ET, 1 October 2026. Interfolio: https://apply.interfolio.com/190236

Output is written to a gitignored folder. Letters of recommendation are
submitted directly by recommenders via Interfolio and are not generated here.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parents[1] / "Haverford-College-CS-TT-2027"

AMBER = RGBColor(0xB9, 0x77, 0x0E)
GREEN = RGBColor(0x1E, 0x84, 0x49)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x44, 0x44, 0x44)

NAME = "Debasish Chakroborti (Joy), Ph.D."
EMAIL = "debasish.chakroborti@usask.ca"
PHONE = "306 914 6212"
ADDR = "410 Stensrud Rd #118, Saskatoon, SK"
WEB = "https://joydeba.github.io/"
ORCID = "https://orcid.org/0000-0002-1597-8162"
GITHUB = "https://github.com/joydeba"
SCHOLAR = "https://scholar.google.com/citations?user=VdrsyT8AAAAJ"

DATE_LINE = "August 17, 2026"
INTERFOLIO = "https://apply.interfolio.com/190236"
SEARCH_EMAIL = "hc-compsciencesearch@haverford.edu"


def set_run(run, *, size=10, bold=False, italic=False, color=DARK, font="Georgia"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_horizontal_line(paragraph, color_hex="B9770E"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def setup_doc(margins=0.7):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(margins)
        section.bottom_margin = Inches(margins)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return doc


def header_block(doc, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(NAME)
    set_run(r, size=16, bold=True, color=DARK, font="Calibri")

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(subtitle)
        set_run(r2, size=11, bold=True, color=AMBER, font="Calibri")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    r = contact.add_run(f"{ADDR}  ·  {PHONE}  ·  {EMAIL}")
    set_run(r, size=9, color=MUTED, font="Calibri")

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(6)
    r = links.add_run(f"{WEB}  ·  {GITHUB}  ·  {ORCID}")
    set_run(r, size=9, color=GREEN, font="Calibri")
    add_horizontal_line(links)


def meta_line(doc, text):
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(text)
    set_run(r, size=9, italic=True, color=MUTED, font="Calibri")
    return meta


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run(r, size=11, bold=True, color=AMBER, font="Calibri")
    add_horizontal_line(p, "B9770E")
    return p


def sub_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, size=10, bold=True, color=GREEN, font="Calibri")
    return p


def body_para(doc, text, *, first_indent=False, space_after=6, size=10, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.2)
    if bold_lead:
        r0 = p.add_run(bold_lead)
        set_run(r0, size=size, bold=True, color=DARK, font="Georgia")
        r = p.add_run(text)
        set_run(r, size=size, color=DARK, font="Georgia")
    else:
        r = p.add_run(text)
        set_run(r, size=size, color=DARK, font="Georgia")
    return p


def bullet(doc, text, *, bold_prefix=None, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        set_run(r0, size=size, bold=True, color=DARK, font="Georgia")
        r1 = p.add_run(text)
        set_run(r1, size=size, color=DARK, font="Georgia")
    else:
        r = p.add_run(text)
        set_run(r, size=size, color=DARK, font="Georgia")
    return p


def entry(doc, left, right_title, body=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(3)
    r1 = p.add_run(f"{left}  ")
    set_run(r1, size=9, bold=True, color=GREEN, font="Calibri")
    r2 = p.add_run(right_title)
    set_run(r2, size=10, bold=True, color=DARK, font="Georgia")
    if body:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        p2.paragraph_format.left_indent = Inches(0.15)
        r = p2.add_run(body)
        set_run(r, size=9.5, color=MUTED, font="Georgia")


def write_text(name, content):
    path = OUT / name
    path.write_text(content.strip() + "\n", encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# 1. COVER LETTER
# ---------------------------------------------------------------------------

def write_cover_letter():
    doc = setup_doc(0.8)
    header_block(doc, "Cover Letter — Tenure-Track Assistant Professor of Computer Science")
    meta_line(doc, "Department of Computer Science  ·  Haverford College  ·  Fall 2027 start")

    body_para(doc, DATE_LINE, space_after=8)
    body_para(
        doc,
        "Search Committee\n"
        "Department of Computer Science\n"
        "Haverford College\n"
        "370 Lancaster Avenue, Haverford, PA 19041\n"
        f"{SEARCH_EMAIL}",
        space_after=8,
    )
    body_para(doc, "Dear Members of the Search Committee,", space_after=8)

    body_para(
        doc,
        "I am applying for the tenure-track Assistant Professor position in Computer Science at Haverford "
        "College, to begin in Fall 2027. I hold a Ph.D. in Computer Science from the University of Saskatchewan "
        "(2024), where I continue as a Living Skies Postdoctoral Fellow. I also teach full-time as an Instructor "
        "in the School of Computing and Digital Innovation at Saskatchewan Polytechnic, and I recently served as "
        "Sessional Lecturer for two sections of introductory computer science at the University of Saskatchewan. "
        "I am looking for a department where undergraduate teaching, undergraduate research, and a durable "
        "scholarly agenda belong together. That is the job I want, and it is the job described in this search.",
        first_indent=True,
    )
    body_para(
        doc,
        "My research asks a practical question: how do we keep software trustworthy on the versions people still "
        "run, not only on the newest development branch? Security patches, bug fixes, and quality problems often "
        "have to be moved backward onto older, stable releases — a skilled, error-prone task called backporting. "
        "I study that task empirically and build tools that help people do it more carefully. During my doctorate, "
        "with Professors Chanchal Roy and Kevin Schneider, I developed Baczer, ReBack, BackSlice, and BackTrans, "
        "which combine repository mining, machine learning, and program analysis. The work has appeared in "
        "Automated Software Engineering, the Journal of Software: Evolution and Process, Empirical Software "
        "Engineering, ACM PACMHCI, and at ICSE, ICPC, MSR, and CASCON. I have also built reusable scientific-workflow "
        "software at the Plant Phenotyping and Imaging Research Centre, and I currently advise Cloud Linux Software "
        "on AI-assisted maintenance of production Linux. The research statement describes how this agenda can run "
        "with undergraduates: public data, concrete questions, existing tools students can extend, and thesis "
        "projects that match Haverford’s fall literature-review / spring implementation rhythm.",
        first_indent=True,
    )
    body_para(
        doc,
        "Haverford’s curriculum is a genuine match for how I already teach. I am ready to staff the core "
        "introductory sequence — CMSC H105 and H106, or the accelerated H107 — from the first semester. I have "
        "carried full responsibility for Principles of Computer Science (CMPT 145) at Saskatchewan, including "
        "lectures, labs, assessments, and three teaching assistants, and I taught Structured Programming and "
        "software-engineering courses for four years as a Lecturer at Khulna University. H105’s emphasis on "
        "specification, testing, peer code review, and reasoning about correctness is the intro course I want to "
        "teach, not a service assignment I would accept on the way to something else. I can also contribute to "
        "computing systems (H251) and, over time, to the software side of HERA — the multi-year system majors "
        "construct across upper-level courses. What I would add to the catalog is an upper-level elective the "
        "department does not currently staff: software design and evolution, and a 300-level course in empirical "
        "software engineering / mining software repositories. Those courses would expand the advanced offerings "
        "without duplicating existing strengths in theory (Lindell), compilers and systems (Wonnacott), fair ML "
        "(Friedler), NLP (Grissom), robotics (Nguyen), or computing education (Dougherty).",
        first_indent=True,
    )
    body_para(
        doc,
        "I take the five-course-unit load, including discussion or lab sections and senior-thesis supervision, as "
        "the center of the job. I have mentored an M.Sc. student, Jarin Tasnim, to co-authored papers at ICSE 2023 "
        "and CASCON 2025, supervised three undergraduate summer researchers, and completed more than 250 hours of "
        "formal instructional training (LIFT, GTI Foundation Training for University Teachers, New Instructor "
        "Orientation, and CETL pedagogy modules). I have taught CS majors, BA students who were not majors, and "
        "high-school teachers. The teaching statement discusses that range, including a first-year course that "
        "did not go as I had planned and what I changed.",
        first_indent=True,
    )
    body_para(
        doc,
        "Haverford’s commitment to social justice, ethical engagement, and a collaborative major is not an add-on "
        "to this application. Software maintenance is an ethical problem as well as a technical one: unpatched "
        "stable software is how vulnerabilities persist in systems people depend on. I want students to learn "
        "that connection in class and in thesis work, not only in a dedicated ethics course. I have completed "
        "Reconciliation Education and Four Seasons of Reconciliation training in Canada; I would continue that "
        "habit of learning on Lenape land and in a Quaker college whose students and faculty already take "
        "inclusion seriously.",
        first_indent=True,
    )
    body_para(
        doc,
        "I am submitting a curriculum vitae, a research statement, and a teaching statement. Three confidential "
        "letters of recommendation, including one that addresses my teaching, will arrive through Interfolio. I "
        "would be glad to talk with the committee about staffing 105/106, adding a software-evolution elective, "
        "and building a small undergraduate research group around maintaining software that people still use.",
        first_indent=True,
    )
    body_para(doc, "Sincerely,", space_after=18)
    body_para(doc, "Debasish Chakroborti (Joy), Ph.D.", space_after=2)
    body_para(doc, f"{EMAIL}  ·  {PHONE}  ·  {WEB}", space_after=2)

    path = OUT / "01_CoverLetter.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2. CURRICULUM VITAE
# ---------------------------------------------------------------------------

def write_cv():
    doc = setup_doc(0.65)
    header_block(doc, "Curriculum Vitae")
    meta_line(
        doc,
        "Application: Tenure-Track Assistant Professor of Computer Science · Haverford College · August 2026",
    )

    section_heading(doc, "Research Interests")
    body_para(
        doc,
        "Software engineering and software evolution — AI-assisted maintenance; automated backporting and "
        "release engineering; technical debt and quality in stable releases; mining software repositories; "
        "empirical studies of developer tools, including LLM- and agent-based workflows; scientific workflow "
        "and data reusability. The program is designed so that undergraduates can join it: public data, "
        "concrete questions, and tools students can extend in senior thesis and summer research.",
        size=9.5,
    )

    section_heading(doc, "Education")
    entry(
        doc,
        "2020–2024",
        "Ph.D., Computer Science / Software Engineering — University of Saskatchewan, Canada",
        "Thesis: Automated Backporting for Efficient Version Management in Software Repositories. "
        "Advisors: Prof. Chanchal K. Roy and Prof. Kevin A. Schneider. Course mark 89%. Awards include Carl McCrosky "
        "Innovation Scholarship (2023), GSA Mark Kroeker Exceptional Student Leadership Award (2023), Ph.D. Citizenship "
        "Award (2024), Education Graduate Award (Friends Club of Saskatoon, 2024), Faculty/Departmental Scholarship, "
        "best poster awards (SOAR 2021, SEMLA 2022, ICSAC 2022), and Graduate Travel Awards (Seattle, Melbourne, Montréal).",
    )
    entry(
        doc,
        "2017–2019",
        "M.Sc., Computer Science / Big Data — University of Saskatchewan, Canada",
        "Thesis: An Intermediate Data-driven Methodology for Scientific Workflow Management System to Support Reusability. "
        "Advisors: Prof. Chanchal K. Roy and Prof. Kevin A. Schneider. Course mark 89%. Saskatchewan Innovation and "
        "Opportunity Scholarship; University of Saskatchewan Faculty Stipend.",
    )
    entry(
        doc,
        "2008–2012",
        "B.Sc., Computer Science and Engineering — Khulna University, Bangladesh",
        "Thesis: Web Service Performance Enhancement for Portable Devices Modifying SOAP Security Principle "
        "(Supervisor: Dr. Masud Rahman). CGPA 3.73/4.00 (1st class, 3rd position). Innovation Award; Dean's Merit / "
        "Merit List Scholarships (2008–2012); Honorable Mention, AUST Inter-University Programming Contest (2009).",
    )

    section_heading(doc, "Academic and Professional Appointments")
    entry(
        doc,
        "2024–Present",
        "Living Skies Postdoctoral Fellow, Department of Computer Science, University of Saskatchewan",
        "NSERC CREATE Software Analytics Research (SOAR). Automated software backporting and AI-driven development "
        "tools; contributing researcher on an NSERC Alliance Advantage industry–academia project (Software Research "
        "Lab, SRLab). Mentor of graduate and undergraduate researchers.",
    )
    entry(
        doc,
        "2024–Present",
        "Instructor, School of Computing and Digital Innovation, Saskatchewan Polytechnic (Regina Campus)",
        "Computer Systems Technology and Cloud Computing & Blockchain programs; designed/rewrote course and lab "
        "materials for most courses taught. Selected faculty participant, China Henan (ZJTIE) Transnational "
        "Education Project (2026).",
    )
    entry(
        doc,
        "2025–2026",
        "Consultant (AI Automation Team), Cloud Linux Software, Inc., USA",
        "Research advisory on AI-based automated backporting tooling for production Linux maintenance.",
    )
    entry(
        doc,
        "2025",
        "Sessional Lecturer, Department of Computer Science, University of Saskatchewan",
        "CMPT 145.3 Principles of Computer Science (Spring/Summer 2025, sections MT1 and MT2; on-campus). Full course "
        "responsibility including deferred/supplemental examinations; guided 3 teaching assistants.",
    )
    entry(
        doc,
        "2024",
        "Software Engineer (Intern), Push Interactions, Saskatoon",
        "Web application enhancement (https://www.wilger.net).",
    )
    entry(
        doc,
        "2023",
        "Ph.D. Intern (Remote), Synesis IT GmbH, Germany",
        "API planning and big-data databases for Traffic Management Intelligent Transportation Systems.",
    )
    entry(
        doc,
        "2017–2024",
        "Teaching Assistant / Research Assistant, University of Saskatchewan",
        "TA (PSAC 40004) for CMPT 141, 145, 214, 270, 280, 370, and Database Management Systems. RA at P2IRC/GIFS: "
        "built RISP and RISPts (Hadoop/Spark scientific-workflow data reusability) and InsCount (UNet-based imaging). "
        "Built SE research tools Baczer, ReBack (CNN), BackTrans (LLM), and BackSlice.",
    )
    entry(
        doc,
        "2013–2017",
        "Lecturer, CSE Discipline, Khulna University, Bangladesh (2017–2024 on study leave)",
        "Lectures, labs, assignments, and examinations; undergraduate thesis/project guidance; developed Khulna "
        "University, ICCIT, and Convocation 2016 registration websites.",
    )
    entry(
        doc,
        "2014–2016",
        "Part-time Lecturer, North Western University; Instructor (Part-time), HSTTI, Khulna",
        "Computer Science courses (BA program); ICT training courses for high-school teachers.",
    )
    entry(
        doc,
        "2012–2013",
        "Software Engineer — Divine IT Limited; Nascenia Limited, Bangladesh",
        "ERP systems (Django/SQL); SnapKnot wedding-photography platform (Ruby on Rails / ASP.NET).",
    )

    section_heading(doc, "Publications")
    body_para(
        doc,
        "24 research items (2015–2026) including 6 journal articles, 10 conference papers, 4 manuscripts in "
        "preparation, and 4 book chapters/awarded posters. Venues include ICSE, ICPC, MSR, CASCON, IEEE Big Data, "
        "Automated Software Engineering, Empirical Software Engineering, JSME, and ACM PACMHCI (EICS).",
        size=9.5,
    )
    sub_heading(doc, "Journal Articles")
    for t in [
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2024). ReBack: Recommending backports in social coding "
        "environments. Automated Software Engineering, 31(1), 18. https://doi.org/10.1007/s10515-024-00416-1",
        "Chakroborti, D., Nath, S., Schneider, K., & Roy, C. (2022). Release conventions of open-source software: "
        "An exploratory study. Journal of Software: Evolution and Process, e2499. https://doi.org/10.1002/smr.2499",
        "Chakroborti, D., Roy, B., & Nath, S. S. (2021). Designing for recommending intermediate states in a scientific "
        "workflow management system. Proc. ACM Hum.-Comput. Interact., 5(EICS). https://doi.org/10.1145/3457145",
        "Herbold, S., Trautsch, A., Ledel, B., et al., including Chakroborti, D. (2021). A fine-grained data set and "
        "analysis of tangling in bug fixing commits. Empirical Software Engineering.",
        "Bairagi, A. K., Mondal, S., & Chakroborti, D. (2017). Securing Bangla text communication using image "
        "steganography with dynamic substitution in IoT environment. Khulna University Studies, 14.",
        "Chakroborti, D., Roy, C., & Schneider, K. (2024). Code propagation in stable version management: A systematic "
        "literature review. Manuscript under review.",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "Conference Proceedings")
    for t in [
        "Chakroborti, D. (2026). Backporting battles on the branches: Agentic vs human fixes. Submitted to ICSE.",
        "Tasnim, J., Chakroborti, D., Roy, C., & Schneider, K. (2025). An insight into the technical debt–fix "
        "trade-off in software backporting. CASCON 2025, Toronto, ON, Canada: IEEE.",
        "Chakroborti, D., Roy, C. K., & Schneider, K. A. (2024). A study of backporting code in open-source software "
        "for characterizing changesets. ICSE-Companion '24. https://doi.org/10.1145/3639478.3643079",
        "Chakroborti, D., Roy, C., & Schneider, K. (2024). BackSlice: Achieving accurate and essential propagation "
        "through backport slicing. Submitted.",
        "Tasnim, J., Chakroborti, D., Roy, C. K., & Schneider, K. A. (2023). How does quality deviate in stable "
        "releases by backporting? ICSE 2023, Melbourne, Australia.",
        "Chakroborti, D., Schneider, K. A., & Roy, C. K. (2022). Backports: Change types, challenges and strategies. "
        "ICPC 2022. https://doi.org/10.1145/3524610.3527920",
        "Bhattacharjee, A., Nath, S. S., Zhou, S., Chakroborti, D., Roy, B., Roy, C. K., & Schneider, K. (2020). An "
        "exploratory study to find motives behind cross-platform forks from Software Heritage dataset. MSR 2020. "
        "https://doi.org/10.1145/3379597.3387512",
        "Chakroborti, D., Mondal, M., Roy, B., Roy, C. K., & Schneider, K. A. (2018). Optimized storing of workflow "
        "outputs through mining association rules. IEEE Big Data 2018. https://doi.org/10.1109/BigData.2018.8622351",
        "Chakroborti, D., & Nath, S. S. (2017). Web service performance enhancement for portable devices modifying "
        "SOAP security principle. ICCIT 2017. https://doi.org/10.1109/ICCITECHN.2017.8281799",
        "Bairagi, A. K., & Chakroborti, D. (2015). Trust based D2D communications for accessing services in Internet "
        "of Things. ICCIT 2015. https://doi.org/10.1109/ICCITechn.2015.7488041",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "In Preparation")
    for t in [
        "Chakroborti, D. An insight into backporting changes: An empirical study with pull-based development.",
        "Chakroborti, D. BranchBench and BranchBridge: Patch-level evaluation for CVE backport adaptation.",
        "Chakroborti, D. Beyond correctness: Technical debt and carbon footprint of LLM-agent versus human backports.",
        "Tasnim, J., Chakroborti, D., Roy, C. K., & Schneider, K. A. Impact of backporting on the quality of stable "
        "software releases: A comparative analysis.",
    ]:
        bullet(doc, t, size=9)

    sub_heading(doc, "Book Chapter and Awarded Posters")
    for t in [
        "Chakroborti, D., Roy, B., Mondal, A., Mostaeen, G., Roy, C. K., Schneider, K. A., & Deters, R. (2020). A data "
        "management scheme for micro-level modular computation-intensive programs in big data platforms. In Data "
        "Management and Analysis (Springer). https://doi.org/10.1007/978-3-030-32587-9_9",
        "BEST POSTER (3rd), SEMLA 2022 — Backporting for version management: Automated changesets integration in "
        "software repositories.",
        "BEST POSTER (2nd), ICSAC 2022 — Detecting and visualizing quality aspects in backporting pull requests "
        "(with mentee J. Tasnim).",
        "BEST POSTER, SOAR Symposium 2021 — Identifying and integrating changesets in backporting.",
    ]:
        bullet(doc, t, size=9)
    body_para(doc, f"Full list: {WEB}research  ·  Google Scholar: {SCHOLAR}", size=9)

    section_heading(doc, "Research Tools & Artifacts")
    bullet(doc, "Baczer — backport changeset identification and analysis at ecosystem scale.", size=9)
    bullet(doc, "ReBack — CNN-based recommendation of backports in social coding environments (ASE journal, 2024).", size=9)
    bullet(doc, "BackSlice — program-slicing approach for accurate, essential backport propagation.", size=9)
    bullet(doc, "BackTrans — LLM-based adaptation/translation of patches across versions and branches.", size=9)
    bullet(doc, "RISP / RISPts — intermediate-data reusability in scientific workflow management systems (Hadoop/Spark, FAIR).", size=9)
    bullet(doc, "InsCount — UNet-based image analysis tool for plant phenotyping (P2IRC/GIFS).", size=9)

    section_heading(doc, "Student Supervision & Mentoring")
    bullet(
        doc,
        "Jarin Tasnim (M.Sc., USask) — software backporting and quality analysis; co-authored ICSE 2023 and "
        "CASCON 2025 papers and an awarded ICSAC poster.",
        bold_prefix="M.Sc. mentoring — ",
        size=9,
    )
    bullet(doc, "Supervised 3 undergraduate summer research students.", bold_prefix="Undergraduate research — ", size=9)
    bullet(doc, "Guided 3 TAs as Sessional Lecturer for CMPT 145 (USask, Spring/Summer 2025).", bold_prefix="Teaching assistants — ", size=9)

    section_heading(doc, "Teaching Experience")
    entry(
        doc,
        "2024–Present",
        "Instructor, Saskatchewan Polytechnic — Computer Systems Technology; Cloud Computing & Blockchain",
        "COOS 190 Server Administration; COOS 294 Cloud Infrastructure Administration; COOS 295 Systems Administration 2; "
        "COET 295 Emerging Technologies; CNET 184 Data Communications and Networking; CWEB 280 Internet Programming and "
        "Web Applications 2; CMPT 145 Principles of Computer Science; CCMP 600–606; DEVP 600 DevOps; PROJ 611 industry "
        "cloud-adoption capstone. Designed/rewrote course and lab materials for most offerings; authored course outlines "
        "for CNET 184 and CWEB 280 that received program approval.",
    )
    entry(
        doc,
        "2025",
        "Sessional Lecturer, University of Saskatchewan — CMPT 145.3 (two sections)",
        "Full course responsibility; SLEQ end-of-course report available (response rate below release threshold for "
        "closed-ended scores; qualitative comments on file).",
    )
    entry(
        doc,
        "2017–2024",
        "Teaching Assistant, University of Saskatchewan",
        "CMPT 141, 145, 214, 270, 280, 370, Database Management Systems — tutorials, labs, marking, help desk, invigilation.",
    )
    entry(
        doc,
        "2013–2017",
        "Lecturer, Khulna University — CSE Discipline",
        "Structured Programming; Software Development Project; Advanced Programming Laboratory; Software Engineering "
        "and Information System (lectures, labs, assignments, examinations).",
    )
    entry(
        doc,
        "2014–2016",
        "Part-time teaching — North Western University (BA CS courses); HSTTI, Khulna (ICT for high-school teachers)",
        None,
    )

    section_heading(doc, "Instructional Development")
    bullet(doc, "LIFT Certificate, Saskatchewan Polytechnic — 90 hours, completed May 2026, with instructional ePortfolio.", size=9)
    bullet(doc, "New Instructor Orientation, Saskatchewan Polytechnic — 18 hours (2025).", size=9)
    bullet(doc, "GTI Foundation Training for University Teachers — 150 hours (2014); CETL/IQAC pedagogy modules (2014–2016).", size=9)
    bullet(doc, "Reconciliation Education; Four Seasons of Reconciliation; TCPS 2: CORE 2022 (research ethics).", size=9)
    bullet(doc, "Selected faculty participant, China Henan (ZJTIE) transnational education project (2026).", size=9)

    section_heading(doc, "Academic Service & Peer Review")
    bullet(doc, "PC Member, Mining Software Repositories (MSR 2026) — Data/Challenge track.", size=9)
    bullet(doc, "PC Member, SOAR Symposium 2025; Lead Organizer, 7th ICSAC / Research Fest.", size=9)
    bullet(doc, "Judge: Research Fest 2025; USask Images of Research (2024); GSA Elevator Pitch (2024); GSA Research Conference (2023).", size=9)
    bullet(doc, "Journals: Journal of Software: Evolution and Process (2023–2025); Journal of Systems and Software (2020, 2022, 2025, 2026).", size=9)
    bullet(doc, "Conferences: ASE 2026/2024/2020; FSE 2026; ICPC 2026; ICSE 2026 SEIS; CHI 2025; SANER 2025/2023; CASCON 2024; ICSME 2023; ISEC 2022; ICSE 2020 Tool Demo; IWSC 2020; ICSA 2020.", size=9)
    bullet(doc, "President, Computer Science Graduate Council, USask (2022–2023); GSA representative on search and planning committees.", size=9)
    bullet(doc, "Volunteer Lead, Digitized Technology — computing outreach for high-school students.", size=9)

    section_heading(doc, "Awards (selected)")
    bullet(doc, "Ph.D. Citizenship Award, University of Saskatchewan (2024); Education Graduate Award, Friends Club of Saskatoon (2024).", size=9)
    bullet(doc, "Carl McCrosky Innovation Scholarship (2023); GSA Mark Kroeker Exceptional Student Leadership Award (2023).", size=9)
    bullet(doc, "Best Poster Awards: SOAR 2021; SEMLA 2022 (3rd); ICSAC 2022 (2nd, with mentee).", size=9)
    bullet(doc, "Saskatchewan Innovation and Opportunity Scholarship; Departmental Scholarship; Ivan and Margaret Toutloff Award.", size=9)
    bullet(doc, "Innovation Award and Dean's Merit Scholarships, Khulna University (2008–2012).", size=9)

    section_heading(doc, "Technical Skills")
    body_para(
        doc,
        "Languages: Python, Java, C/C++, C#, Ruby, JavaScript, R, PHP, MATLAB, SQL, LaTeX. Methods: mining software "
        "repositories, empirical software engineering, ML/LLM pipelines for code, program analysis and slicing. "
        "Human languages: English, Bengali.",
        size=9,
    )

    section_heading(doc, "References")
    body_para(
        doc,
        "Three confidential letters of recommendation are being submitted directly via Interfolio, including at "
        "least one that addresses teaching, as required by the posting.",
        size=9.5,
    )

    path = OUT / "02_CV.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 3. RESEARCH STATEMENT
# ---------------------------------------------------------------------------

def write_research_statement():
    doc = setup_doc(0.8)
    header_block(doc, "Research Statement")
    meta_line(
        doc,
        "Tenure-Track Assistant Professor of Computer Science · Haverford College · Fall 2027\n"
        "Begins with a general-audience summary, as requested by the search committee",
    )

    section_heading(doc, "Summary (for a general audience)")
    body_para(
        doc,
        "Most of the software that runs phones, hospitals, and public infrastructure is not the newest version. "
        "It is an older, “stable” version that still has to receive security fixes and bug repairs. Moving a fix "
        "from new code onto that older version is called backporting. It is skilled work, and it is easy to get "
        "wrong: a patch that is harmless on the newest branch can break an older one, or leave a hole unfixed. "
        "My research studies how this work is actually done in large open-source projects, and builds methods and "
        "tools that help people do it more carefully — which changes should be copied back, which parts of a change "
        "are essential, and when an automated suggestion would make the older software worse rather than better. "
        "Undergraduates can join this work. The data are public. The questions are concrete. The tools already "
        "exist for students to extend. A senior thesis can be a careful reading of how one project handles stable "
        "releases, or a small experiment that tests a new idea on real repositories. That is the research life I "
        "want at a college whose students write, present, and think hard in a required thesis.",
        first_indent=True,
    )

    section_heading(doc, "1. What I work on, and why it belongs at Haverford")
    body_para(
        doc,
        "I work in empirical software engineering, with a long-running focus on software evolution: how long-lived "
        "systems change, how fixes propagate across branches, and how quality and technical debt accumulate on the "
        "releases people actually use. The scholarly claim is not that automation should replace maintainers. It is "
        "that backporting and related maintenance tasks are regular enough, and consequential enough, to be studied "
        "with the same care we give to compilers or algorithms — with measurements, artifacts, and arguments other "
        "people can check.",
        first_indent=True,
    )
    body_para(
        doc,
        "A small undergraduate college is a good home for this agenda, not a compromise. I do not need a large "
        "graduate laboratory or a private industrial dataset. I need students who will read papers closely, write "
        "clearly, run careful studies, and stay with a question for a year. That is already how Haverford’s senior "
        "thesis is designed: a literature review in the fall, and, for students who continue, an implementation or "
        "deeper analysis in the spring. My existing tools (Baczer, ReBack, BackSlice, BackTrans) and public "
        "repository data give those students something real to join, rather than a waiting period while I recruit "
        "Ph.D. students. Industry contact with Cloud Linux, which maintains long-lived Linux distributions, can "
        "supply well-scoped undergraduate problems about stable branches without turning the lab into a consulting "
        "shop.",
        first_indent=True,
    )
    body_para(
        doc,
        "The work also sits in an open space in the current department. Haverford’s tenure-track faculty cover "
        "theory, compilers and systems, fair machine learning, natural language processing, robotics, and computing "
        "education. Empirical software engineering and software evolution are not represented. I would not duplicate "
        "those programs. I would give students a way to study the software they themselves write — including HERA, "
        "the multi-year system majors construct — as something that has to be designed, tested, and kept alive, not "
        "only completed for a deadline.",
        first_indent=True,
    )

    section_heading(doc, "2. Past and present scholarship")
    body_para(
        doc,
        "A pipeline for studying and assisting backporting. In my Ph.D. at the University of Saskatchewan, advised "
        "by Professors Chanchal K. Roy and Kevin A. Schneider, I built a sequence of systems around one maintenance "
        "task. Baczer characterizes backporting changesets at ecosystem scale (ICPC 2022; ICSE Companion 2024). "
        "ReBack uses convolutional networks over social-coding signals — pull requests, reviews, discussion — to "
        "recommend which changes should be backported (Automated Software Engineering, 2024). BackSlice uses "
        "program slicing so that only the essential parts of a change are propagated. BackTrans uses large language "
        "models to adapt patches across diverging branches. Alongside the tools, I studied how open-source projects "
        "actually cut releases (Journal of Software: Evolution and Process, 2022) and contributed to a widely used "
        "dataset of tangled bug-fixing commits (Empirical Software Engineering, 2021).",
        first_indent=True,
    )
    body_para(
        doc,
        "Quality, technical debt, and students as co-authors. With M.Sc. student Jarin Tasnim I measured how quality "
        "shifts in stable releases after backporting (ICSE 2023) and how maintainers trade a fix against the debt "
        "they take on (CASCON 2025). Her ICSAC poster on visualizing quality in backporting pull requests received "
        "a best-poster award. That collaboration is the mentoring pattern I would bring to Haverford theses: a "
        "question small enough to finish, data the student can defend, writing early, and a public presentation. I "
        "have also supervised three undergraduate summer researchers on related projects.",
        first_indent=True,
    )
    body_para(
        doc,
        "Current work. As a Living Skies Postdoctoral Fellow I am asking a question the field cannot avoid: when an "
        "LLM agent proposes a backport, is the result merely “correct,” or does it leave behind debt, missing tests, "
        "or energy cost that a human maintainer would not have accepted? A 2026 ICSE submission compares agentic "
        "and human fixes head to head. In-preparation work (BranchBench / BranchBridge) builds patch-level "
        "benchmarks for CVE backport adaptation. I am also finishing a systematic literature review on code "
        "propagation in stable version management. The point of this phase is not to chase every new model. It is "
        "to keep evaluation honest as the tools change.",
        first_indent=True,
    )
    body_para(
        doc,
        "Earlier applied systems. During my M.Sc. and as a research assistant at the Plant Phenotyping and Imaging "
        "Research Centre / Global Institute for Food Security, I built RISP and RISPts so that intermediate data in "
        "scientific workflows could be found and reused (IEEE Big Data 2018; ACM PACMHCI 2021; Springer chapter), "
        "and InsCount, an imaging tool for phenotyping. That work taught me to ship research software that other "
        "scientists can run — a useful habit when undergraduates need a system they can clone on day one.",
        first_indent=True,
    )

    section_heading(doc, "3. Planned projects, with undergraduates in the work")
    body_para(
        doc,
        "I would keep a small, steady group: typically one or two senior-thesis students a year, a summer researcher "
        "when funding allows, and independent-study students who want a shorter project. Every project should produce "
        "writing a second reader can assess, and, when it earns it, an artifact or a student-authored workshop or "
        "poster. The three themes below are sized for that group, not for a graduate pipeline.",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 1 — How stable software actually changes. Students can replicate and extend my empirical studies on "
        "new projects or new years of data: what gets backported, what fails, and how quality metrics move. A "
        "fall thesis can be a literature review plus a careful reproduction; a spring continuation can add a new "
        "comparison (for example, security patches versus feature backports). This theme trains the skills "
        "Haverford already lists as thesis goals: separating a problem from a solution, describing a method, and "
        "comparing alternatives.",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 2 — Tools students can read, break, and improve. Baczer, ReBack, BackSlice, and BackTrans are not "
        "finished products. An undergraduate can add a baseline, test a slicing heuristic on a new language, or "
        "build a small visualization that helps a person inspect a recommended backport. BranchBench is deliberately "
        "a benchmark project: students can contribute cases, documentation, and evaluation scripts. I will not "
        "hand a first-year researcher a training run they cannot explain. I will hand them a question whose "
        "answer they can defend in a poster session.",
        first_indent=True,
    )
    body_para(
        doc,
        "Theme 3 — Maintenance as an ethical and civic problem. Unpatched stable software is how known "
        "vulnerabilities remain in systems people depend on. That claim connects to work already on campus in "
        "fairness, accountability, and the social life of algorithms, without copying it. A thesis here might "
        "study how projects communicate risk when a fix cannot be backported, or how student teams on HERA handle "
        "change after the original authors have graduated. I would welcome conversation with colleagues in CS and "
        "in the wider college about where that work should live — in a software-evolution elective, in thesis "
        "supervision, or in a module inside the intro sequence when we talk about testing and review.",
        first_indent=True,
    )
    body_para(
        doc,
        "Year-one starter projects (examples a junior could take): (1) a reproduction study of release conventions "
        "in one well-documented open-source project; (2) a documented extension of BackSlice with tests and a "
        "negative-result write-up if the heuristic fails; (3) a small human-facing study of how students reason "
        "about a backport recommendation, suitable for a thesis that is more exposition than new systems work. "
        "I would advertise these in the senior seminar’s advisor-selection process and as summer possibilities.",
        first_indent=True,
    )

    section_heading(doc, "4. How the program is sustained")
    body_para(
        doc,
        "At Haverford I would seek support that fits undergraduate research rather than a graduate-lab budget: "
        "internal KINSC and faculty research grants; Tri-Co and Bi-Co collaboration where interests overlap; NSF "
        "CRII in the first eligible window, with a later CAREER only if the undergraduate program is actually "
        "running; and REU-style summer support when it does not distort the teaching calendar. My work is "
        "compute-moderate — repository mining and modest model use, not frontier training — so it can live on "
        "shared college resources and student laptops. Open artifacts remain a requirement: datasets, benchmarks, "
        "and tools released with enough documentation that a Haverford student can pick them up after a predecessor "
        "graduates. That is how a research group survives at a college without graduate students.",
        first_indent=True,
    )
    body_para(
        doc,
        "I am not proposing to transplant a Canadian funding plan onto a U.S. liberal-arts department. I am "
        "proposing a scholarly life that is already compatible with one: a focused question, students who write, "
        "and software that other people can inspect. If that work also speaks to colleagues at Bryn Mawr or "
        "Swarthmore, or to scientific-computing concentrators who maintain research code, I will meet them there. "
        "The center of gravity stays in Haverford’s CS major, the thesis, and the intro sequence that feeds both.",
        first_indent=True,
    )

    path = OUT / "03_ResearchStatement.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4. TEACHING STATEMENT
# ---------------------------------------------------------------------------

def write_teaching_statement():
    doc = setup_doc(0.8)
    header_block(doc, "Teaching Statement")
    meta_line(
        doc,
        "Tenure-Track Assistant Professor of Computer Science · Haverford College · Fall 2027\n"
        "Prior experience · diverse students · varied pedagogies · successes and challenges",
    )

    section_heading(doc, "1. What I am trying to do in a classroom")
    body_para(
        doc,
        "I want students to leave a course able to specify a problem, write a program they can defend, and explain "
        "why they believe it is right — by testing, by reading someone else’s code, and, when the material calls "
        "for it, by a short argument. That aim is close to Haverford’s own learning goals for computer science: "
        "abstraction, correctness, complexity, and clear communication, including the ethical implications of the "
        "work. I do not treat those as slogans. They are the weekly work of an intro lab and the year-long work "
        "of a thesis.",
        first_indent=True,
    )
    body_para(
        doc,
        "Three habits keep me honest. I publish what success looks like (outcomes, rubrics, weekly targets) so "
        "students are not guessing at hidden rules. I put students in contact with real artifacts — a failing test, "
        "a messy pull request, a design that has to change — rather than only toy exercises. And I change the "
        "course when the room tells me it is not working, which I discuss below because this committee asked for "
        "challenges as well as successes.",
        first_indent=True,
    )

    section_heading(doc, "2. Prior teaching experience")
    body_para(
        doc,
        "Introductory computer science. In Spring/Summer 2025 I was Sessional Lecturer for two on-campus sections "
        "of CMPT 145, Principles of Computer Science, at the University of Saskatchewan. I had full responsibility "
        "for lectures, labs, assessments, deferred and supplemental exams, and three teaching assistants. I have "
        "also taught the same introductory level as a TA (CMPT 141 and 145) and currently teach a Principles of "
        "Computer Science offering at Saskatchewan Polytechnic. That is the closest analogue I have to CMSC H105 "
        "and H106, including the weekly lab. H105’s attention to specification, testing, peer code review, and "
        "both functional and imperative styles is work I already do; I would be glad to teach it in Haverford’s "
        "idiom rather than import a different intro.",
        first_indent=True,
    )
    body_para(
        doc,
        "Programming, software engineering, and projects. As Lecturer in CSE at Khulna University (2013–2017) I "
        "taught Structured Programming, Advanced Programming Laboratory, Software Development Project, and Software "
        "Engineering and Information System. Those years included undergraduate project guidance and the ordinary "
        "work of setting exams, running labs, and sitting with students who were stuck. That is the background I "
        "would bring to a new upper-level elective in software design and evolution, and to helping majors reason "
        "about HERA as a system that has to be built and then kept.",
        first_indent=True,
    )
    body_para(
        doc,
        "A broader teaching life. At Saskatchewan Polytechnic I teach across systems administration, networking, "
        "web applications, cloud, and project courses, and I have rewritten labs and outlines for most of what I "
        "teach, including two outlines that received formal program approval. I do not propose to turn Haverford "
        "into a vocational program. I mention this work because it is where I learned to write labs that a mixed "
        "room can finish, and because applied electives (for example, a carefully scoped course on software "
        "construction) can expand the catalog without chasing short-term fashions. I have also taught Computer "
        "Science courses for BA students at North Western University who were not CS majors, and ICT courses for "
        "high-school teachers at HSTTI in Khulna. Those rooms changed how I explain things. They are part of why "
        "I am not afraid of Haverford’s mix of intended majors, minors, and students meeting CS for the first time.",
        first_indent=True,
    )
    body_para(
        doc,
        "I have completed more than 250 hours of formal instructional training: LIFT (90 hours, 2026), New "
        "Instructor Orientation (18 hours, 2025), GTI Foundation Training for University Teachers (150 hours, 2014), "
        "and CETL/IQAC pedagogy modules at Khulna University. Training is not the same as being a good teacher. It "
        "is evidence that I treat teaching as work I study, not as a talent I claim.",
        first_indent=True,
    )

    section_heading(doc, "3. Students from different backgrounds")
    body_para(
        doc,
        "I have taught in Bangladesh and in Canada, in a public university, a polytechnic, a BA program, and a "
        "teachers’ training institute. The constant is not a single classroom culture. It is that students arrive "
        "with unequal preparation, unequal permission to speak, and unequal experience of being told they belong in "
        "computing. I try to reduce the hidden curriculum: weekly targets, worked examples, rubrics, and a rule "
        "that asking a question is a contribution. In team projects I assign roles and check participation so that "
        "the person who already knows git does not silently become the only author.",
        first_indent=True,
    )
    body_para(
        doc,
        "During COVID I coordinated student support work through the Graduate Students’ Association. As CSGC "
        "president I learned how often academic trouble is logistical and social, not a lack of ability. I have "
        "completed Reconciliation Education and Four Seasons of Reconciliation, which in my Canadian classrooms "
        "meant taking land, history, and respect as part of professional practice rather than a first-day slide. "
        "At Haverford I would not paste that Canadian training onto Lenape land and call it done. I would keep the "
        "habit: learn the place, listen to colleagues and students, and design courses so that students who are "
        "not already fluent in the culture of CS can still do rigorous work. Haverford’s own language of diversity, "
        "inclusion, and ethical engagement is the standard I would be accountable to.",
        first_indent=True,
    )

    section_heading(doc, "4. Varied pedagogies")
    body_para(
        doc,
        "I lecture, but I do not only lecture. In intro courses I use live coding with planted mistakes and ask "
        "the room to catch them; think-pair-share on a specification before anyone opens an editor; and peer code "
        "review, which H105 already names as a technique I would be joining, not inventing. Labs are sequenced so "
        "that the first tasks are finishable in the room and later tasks require a design choice the student has "
        "to justify. In software-engineering and project courses I use milestone demos, code walkthroughs, and "
        "short postmortems after a deliverable, because students learn as much from explaining a failure as from "
        "a green test suite.",
        first_indent=True,
    )
    body_para(
        doc,
        "On AI tools I have a policy rather than a panic. Students may use assistants if the course says so, but "
        "they must be able to explain, test, and take responsibility for what they submit. In an intro course that "
        "often means forbidding generated solutions on early labs and then, later, using a generated fragment as "
        "an object of critique. In an upper-level software-evolution course it means measuring what an automated "
        "suggestion actually does to a stable branch — which is also my research. I want students to leave able to "
        "use new tools without surrendering the ability to tell whether the tools are any good.",
        first_indent=True,
    )

    section_heading(doc, "5. Successes and a challenge I take seriously")
    body_para(
        doc,
        "What has gone well: students in my programming and project courses have shipped working systems through "
        "staged deadlines; I have rewritten labs when queues at the help desk showed a bottleneck; a mentored "
        "graduate student became a co-author at ICSE and CASCON; and I was selected to teach abroad in "
        "Saskatchewan Polytechnic’s China Henan partnership, which is a vote of confidence in how I run a lab. "
        "Qualitative comments from CMPT 145 mentioned clear explanations for early-stage learners and responsive "
        "communication. I am glad of those comments. I do not treat them as a complete picture.",
        first_indent=True,
    )
    body_para(
        doc,
        "What did not go as I had planned: CMPT 145 in summer 2025. I met the problem every intro instructor meets "
        "and that Haverford’s 105/106 sequence is built to handle — students in the same room with very different "
        "preparation. Some had written substantial programs. Others were meeting the language for the first time "
        "and did not want to look lost. My first weeks over-explained for one group and still moved too quickly "
        "for the other. I changed three things. I published weekly learning targets so “caught up” was visible. I "
        "used live coding with deliberate mistakes to give quieter students a structured way to speak. I asked TAs "
        "to log which lab steps caused the longest queues, and I rewrote those labs before the following week. "
        "Separately, only three of thirty-eight students completed the official end-of-course questionnaire, so "
        "closed-ended scores were not released. I read that low response as a failure of how I closed the course — "
        "I did not treat evaluation as part of teaching — and I now build a last-week feedback practice into the "
        "plan rather than hoping a survey email will do the work. I would rather say that here than offer a "
        "statement in which every class succeeds.",
        first_indent=True,
    )

    section_heading(doc, "6. What I would teach at Haverford")
    body_para(
        doc,
        "Ready in the first year. CMSC H105 and H106 (and H107 if the department wants the accelerated intro). I "
        "treat staffing the introductory sequence as the primary teaching contribution of this position, consistent "
        "with the advertisement. I can also contribute discussion or lab sections, which the five-unit load counts, "
        "and I am willing to learn HERA well enough to support it in systems-adjacent teaching.",
        first_indent=True,
    )
    body_para(
        doc,
        "New electives that expand the advanced list. (1) Software Design and Evolution — specification, testing, "
        "design for change, and the life of a system after the first release; natural pairing with HERA and with "
        "senior projects. (2) Empirical Software Engineering / Mining Software Repositories — a 300-level course "
        "in which students replicate a published study, handle real repository data, and write. Either course could "
        "serve as an elective in the major; I would discuss with the department whether the 300-level offering is "
        "a better fit in the applications track or as a free elective. I am not proposing a cloud or blockchain "
        "curriculum. Those are things I know how to teach. They are not what this catalog is for.",
        first_indent=True,
    )
    body_para(
        doc,
        "Senior thesis and independent study. I will take thesis students. Fall work can be a literature review "
        "that a second reader can assess at the January checkpoint; spring work can be a reproduction, a small "
        "tool, or original exposition. Not every thesis needs to be a publication. Every thesis needs to show that "
        "the student can think past the textbook. That standard is already in the department’s senior-project "
        "learning goals, and I would hold my students to it.",
        first_indent=True,
    )
    body_para(
        doc,
        "I am applying because I want to teach computer science as a liberal art: a discipline in which students "
        "learn to think, write, and take responsibility for systems other people will have to live with. Haverford "
        "is a place that already describes the major that way. I would be honored to help staff it.",
        first_indent=True,
    )

    path = OUT / "04_TeachingStatement.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 5. LETTER-WRITER BRIEFING (not submitted)
# ---------------------------------------------------------------------------

def write_letter_briefing():
    doc = setup_doc(0.8)
    header_block(doc, "Letter-Writer Briefing (internal — do not submit)")
    meta_line(
        doc,
        "Haverford College · Tenure-Track Assistant Professor of Computer Science · Interfolio 190236\n"
        "Deadline for all materials, including letters: 11:59pm ET, Thursday, 1 October 2026",
    )

    section_heading(doc, "What the committee requires")
    body_para(
        doc,
        "Three confidential letters, submitted directly by recommenders through Interfolio. At least one letter "
        "must address skill as a teacher. The search is a liberal-arts, undergraduate-only department (Haverford "
        "College, outside Philadelphia). The committee has said it wants to hear from the candidate in their own "
        "words and will look poorly on AI-generated application materials; letters should likewise be specific.",
        first_indent=True,
    )
    bullet(doc, "https://apply.interfolio.com/190236", bold_prefix="Interfolio: ", size=9.5)
    bullet(doc, SEARCH_EMAIL, bold_prefix="Search email: ", size=9.5)
    bullet(doc, "Fall 2027; Ph.D. already complete.", bold_prefix="Start: ", size=9.5)

    section_heading(doc, "Suggested slate")
    body_para(
        doc,
        "Do not send three research-only letters. Confirm the teaching letter this week. Banani Roy is a strong "
        "research/collaboration letter; she should not be used as the teaching letter unless she has actually "
        "observed teaching.",
        first_indent=True,
    )
    bullet(
        doc,
        "Prof. Chanchal K. Roy, University of Saskatchewan — chanchal.roy@usask.ca — Ph.D./M.Sc. co-advisor; "
        "current postdoctoral supervisor. Ask him to speak to research independence, mentoring of Jarin Tasnim, "
        "and whether the agenda can run with undergraduates.",
        bold_prefix="Letter 1 (research): ",
        size=9.5,
    )
    bullet(
        doc,
        "Prof. Kevin A. Schneider, University of Saskatchewan — kevin.schneider@usask.ca — Ph.D./M.Sc. co-advisor. "
        "Ask him to speak to scholarly quality, collaboration, and communication.",
        bold_prefix="Letter 2 (research): ",
        size=9.5,
    )
    bullet(
        doc,
        "Someone who has seen you teach. Preferred order: (a) a University of Saskatchewan colleague who observed "
        "CMPT 145 or can speak to sessional teaching; (b) a Saskatchewan Polytechnic supervisor or peer who has "
        "seen labs and curriculum rewrite; (c) a Khulna University CSE colleague who can speak to four years of "
        "lecturing. This letter is the one Interfolio packets often miss. Ask now.",
        bold_prefix="Letter 3 (teaching — required): ",
        size=9.5,
    )

    section_heading(doc, "Talking points to send each writer")
    bullet(doc, "Open-area SLAC search; teaching the intro sequence (CMSC H105/H106) is named in the ad.", size=9.5)
    bullet(doc, "Teaching load is five units, typically three or four regular courses plus labs and senior thesis.", size=9.5)
    bullet(doc, "Research must thrive without graduate students; thesis is required of all majors.", size=9.5)
    bullet(doc, "Software engineering / evolution would expand electives; department currently has no SE faculty.", size=9.5)
    bullet(doc, "Mention specific teaching if they saw it: CMPT 145 (two sections, 2025), Khulna lecturing, Sask Polytech labs.", size=9.5)
    bullet(doc, "Mentee Jarin Tasnim: co-authored ICSE 2023 and CASCON 2025.", size=9.5)

    section_heading(doc, "Draft request (adapt before sending)")
    body_para(
        doc,
        "I am applying for a tenure-track Assistant Professor position in Computer Science at Haverford College "
        "(liberal arts, undergraduate thesis required, start Fall 2027, deadline 1 October 2026). The committee "
        "wants three Interfolio letters; at least one must address teaching. I am asking you to write the "
        "[research / teaching] letter. I will send a CV, the job ad, and a short paragraph on why this college. "
        "Interfolio will email you a confidential upload link. Thank you for considering this — I know the "
        "deadline is close.",
        first_indent=True,
    )

    path = OUT / "05_LetterWriterBriefing_DO_NOT_SUBMIT.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 6. INTERFOLIO CHECKLIST / README
# ---------------------------------------------------------------------------

def write_readme():
    doc = setup_doc(0.8)
    header_block(doc, "Application checklist (internal)")
    meta_line(doc, "Haverford College · CS TT · Interfolio 190236 · Deadline 11:59pm ET, 1 October 2026")

    section_heading(doc, "Submit via Interfolio")
    bullet(doc, "Cover letter — 01_CoverLetter.docx (export PDF if Interfolio prefers PDF).", size=9.5)
    bullet(doc, "CV — 02_CV.docx. You may instead upload cv-academic.pdf from the website if you prefer that layout.", size=9.5)
    bullet(doc, "Research statement — 03_ResearchStatement.docx. Must keep the general-audience opening.", size=9.5)
    bullet(doc, "Teaching statement — 04_TeachingStatement.docx.", size=9.5)
    bullet(doc, "Three letters — recommenders upload themselves. You only enter their emails in Interfolio.", size=9.5)

    section_heading(doc, "Do not submit")
    bullet(doc, "05_LetterWriterBriefing_DO_NOT_SUBMIT.docx", size=9.5)
    bullet(doc, "This checklist.", size=9.5)
    bullet(doc, "A separate EDIA statement (not requested; themes are already in the teaching statement).", size=9.5)
    bullet(doc, "Guelph or other institutional packets. Do not reuse those files.", size=9.5)

    section_heading(doc, "Before you upload")
    body_para(
        doc,
        "The committee wrote that it will look poorly on AI-generated materials. Read every document aloud. Put "
        "back any sentence that does not sound like you. Add one detail only you know (a specific lab, a specific "
        "student question, a specific thesis idea). Change anything that feels generic. Then export to PDF, check "
        "page breaks, and submit. Letters: request them this week.",
        first_indent=True,
    )
    body_para(
        doc,
        f"Portal: {INTERFOLIO}\nQuestions: {SEARCH_EMAIL}\n"
        "Provost page: https://www.haverford.edu/provost/available-positions/tenure-track-search-computer-science",
        space_after=6,
    )

    path = OUT / "00_README_checklist.docx"
    doc.save(path)
    return path


def write_readme_md():
    text = f"""# Haverford College — CS tenure-track application (Fall 2027)

**Do not commit this folder.** It is gitignored.

| Item | File | Submit? |
|---|---|---|
| Cover letter | `01_CoverLetter.docx` | Yes |
| CV | `02_CV.docx` | Yes (or `cv-academic.pdf` from the site) |
| Research statement | `03_ResearchStatement.docx` | Yes |
| Teaching statement | `04_TeachingStatement.docx` | Yes |
| Three letters | Interfolio (recommenders) | Yes — they upload |
| Letter-writer briefing | `05_LetterWriterBriefing_DO_NOT_SUBMIT.docx` | **No** |
| This checklist | `00_README_checklist.docx` / `README.md` | **No** |

- **Deadline:** 11:59pm ET, 1 October 2026
- **Portal:** {INTERFOLIO}
- **Search email:** {SEARCH_EMAIL}
- **Start:** Fall 2027

## Required by the ad

Cover letter, CV, research statement, teaching statement, three confidential letters (at least one on teaching).

The research statement must open with a general-audience summary. The teaching statement should cover prior experience, students from diverse backgrounds, varied pedagogies, and successes **and** challenges. The committee will look poorly on AI-generated materials: read these drafts aloud and revise into your own wording before upload.

## Regenerate

From the repository root:

```
python3 tools/haverford_application_docs.py
```

Convert to PDF if Interfolio asks for PDF (macOS):

```
textutil -convert pdf 01_CoverLetter.docx
```

or open in Word / Pages and export. `textutil` PDF conversion is not always available for `.docx`; Word export is the reliable path.
"""
    return write_text("README.md", text)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [
        write_readme(),
        write_readme_md(),
        write_cover_letter(),
        write_cv(),
        write_research_statement(),
        write_teaching_statement(),
        write_letter_briefing(),
    ]
    for p in paths:
        print(p)


if __name__ == "__main__":
    main()
